# -*- coding: utf-8 -*-
"""对比改动前后：一次 12h 满领（720 tick）玩家实际看到的道具种类/件数。

装备判定：凡出现在 20004~20017 这 13 个品质组里的 itemId 都算装备。
"""
import os
import random
import sys
from collections import defaultdict

import openpyxl

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
BASE = os.path.join(ROOT, "data_execl", "base_data")
BAK = os.path.join(BASE, "_backup_before_tune")

HEADER_ROWS = 6
TICKS = 720
RUNS = 400
LEVELS = [1011, 1012, 1013, 1014]
EQUIP_GROUPS = {20004, 20005, 20006, 20007, 20008, 20009, 20010,
                20012, 20013, 20014, 20015, 20016, 20017}
NAMES = {1: "简单", 2: "普通", 3: "困难", 4: "噩梦"}


def _num(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


def load(path):
    groups, items = {}, {}
    wb = openpyxl.load_workbook(os.path.join(path, "RewardGroup.xlsx"))
    ws = wb.worksheets[0]
    for r in range(HEADER_ROWS + 1, ws.max_row + 1):
        gid = _num(ws.cell(r, 2).value)
        if gid is None:
            continue
        groups[int(gid)] = {
            "rule": int(_num(ws.cell(r, 4).value) or 0),
            "items": [int(_num(x)) for x in str(ws.cell(r, 5).value).split("#") if _num(x) is not None],
            "max": int(_num(ws.cell(r, 6).value) or 0),
        }
    wb = openpyxl.load_workbook(os.path.join(path, "RewardItemConfig.xlsx"))
    ws = wb.worksheets[0]
    for r in range(HEADER_ROWS + 1, ws.max_row + 1):
        rid = _num(ws.cell(r, 2).value)
        iid = _num(ws.cell(r, 3).value)
        if rid is None or iid is None:
            continue
        items[int(rid)] = {
            "iid": int(iid),
            "chance": float(_num(ws.cell(r, 4).value) or 0),
            "lo": int(_num(ws.cell(r, 6).value) or 0),
            "hi": int(_num(ws.cell(r, 7).value) or 0),
        }
    wb = openpyxl.load_workbook(os.path.join(path, "MainLevelConfig.xlsx"))
    ws = wb.worksheets[0]
    hdr = [c.value for c in ws[1]]
    c_id, c_rr = hdr.index("Id") + 1, hdr.index("RandomReward") + 1
    c_diff = hdr.index("Difficulty") + 1 if "Difficulty" in hdr else None
    levels = {}
    for r in range(HEADER_ROWS + 1, ws.max_row + 1):
        lid = _num(ws.cell(r, c_id).value)
        if lid is None:
            continue
        levels[int(lid)] = {
            "gids": [int(_num(x)) for x in str(ws.cell(r, c_rr).value or "").split("#") if _num(x) is not None],
            "diff": int(_num(ws.cell(r, c_diff).value) or 0) if c_diff else 0,
        }
    return groups, items, levels


def equip_ids(groups, items):
    out = set()
    for gid in EQUIP_GROUPS:
        for rid in groups.get(gid, {}).get("items", []):
            it = items.get(rid)
            if it:
                out.add(it["iid"])
    return out


def pool_of(gid, groups, items):
    out = []
    for rid in groups.get(gid, {}).get("items", []):
        it = items.get(rid)
        # 注意：max=0 的「空位」条目必须留在池子里，它代表「这次没掉」，不能过滤掉
        if it and it["chance"] > 0:
            out.append(it)
    return out


def simulate(gids, groups, items, equips, rng):
    pools = []
    for gid in gids:
        p = pool_of(gid, groups, items)
        if p:
            pools.append((groups[gid]["rule"], groups[gid]["max"], p))
    got = defaultdict(int)
    for _ in range(TICKS):
        for rule, maxn, pool in pools:
            if rule == 1:
                picks = pool
            elif rule == 4:
                picks = _weighted(pool, max(1, maxn or 1), rng)
            elif rule == 2:
                picks = _weighted(pool, min(maxn or len(pool), len(pool)), rng)
            else:
                picks = _successive(pool, maxn or len(pool), rng)
            for it in picks:
                q = rng.randint(it["lo"], it["hi"])
                if q > 0:
                    got[it["iid"]] += q
    total = sum(got.values())
    eq = sum(v for k, v in got.items() if k in equips)
    return len(got), eq, total


def _weighted(pool, n, rng):
    pool = list(pool)
    out = []
    for _ in range(min(n, len(pool))):
        tot = sum(p["chance"] for p in pool)
        if tot <= 0:
            break
        r = rng.random() * tot
        acc = 0.0
        for i, p in enumerate(pool):
            acc += p["chance"]
            if r <= acc:
                out.append(pool.pop(i))
                break
    return out


def _successive(pool, n, rng):
    out = []
    for p in pool:
        if len(out) >= n:
            break
        if rng.random() < p["chance"]:
            out.append(p)
    return out


def main():
    rng = random.Random(20261020)
    data = {}
    for tag, path in (("before", BAK), ("after", BASE)):
        g, i, l = load(path)
        data[tag] = (g, i, l, equip_ids(g, i))

    print(f"{'难度':<4}{'关卡':<7}{'组数':<11}{'种类':<15}{'装备件':<15}{'总件':<15}")
    for lv in LEVELS:
        cells = []
        for tag in ("before", "after"):
            g, i, l, eq = data[tag]
            cfg = l.get(lv)
            k = e = t = 0
            for _ in range(RUNS):
                a, b, c = simulate(cfg["gids"], g, i, eq, rng)
                k += a; e += b; t += c
            cells.append((len(cfg["gids"]), k / RUNS, e / RUNS, t / RUNS, cfg["diff"]))
        b, a = cells
        print(f"{NAMES.get(b[4], '?'):<4}{lv:<7}"
              f"{b[0]:>2}->{a[0]:<2}    "
              f"{b[1]:>5.1f}->{a[1]:<5.1f}    "
              f"{b[2]:>5.1f}->{a[2]:<5.1f}    "
              f"{b[3]:>6.1f}->{a[3]:<6.1f}")


if __name__ == "__main__":
    main()
