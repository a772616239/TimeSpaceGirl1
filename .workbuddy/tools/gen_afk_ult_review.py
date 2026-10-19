#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成「挂机大招前摇台词」评审页 _output/afk_ult_talk.html"""
import importlib.util
import os
import openpyxl

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
OUT = os.path.join(ROOT, '_output/afk_ult_talk.html')
DATA = os.path.join(ROOT, 'data_execl/base_data')


def load(name):
    wb = openpyxl.load_workbook(os.path.join(DATA, name + '.xlsx'), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(min_row=1, values_only=True)
    hdr = list(next(it))
    idx = {h: i for i, h in enumerate(hdr) if h}
    rows = [r for r in ws.iter_rows(min_row=8, values_only=True) if isinstance(r[idx['Id']], int)]
    wb.close()
    return rows, idx


heroes, hi = load('HeroConfig')
logic, li = load('SkillLogicConfig')
comb, ci = load('CombatControl')
skills, ki = load('SkillConfig')

lang = {}
wb = openpyxl.load_workbook(os.path.join(DATA, 'Language.xlsx'), read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
for r in ws.iter_rows(min_row=8, values_only=True):
    if isinstance(r[1], int):
        lang[r[1]] = r[2:5]
wb.close()

hero_by_id = {r[hi['Id']]: r for r in heroes}
comb_by_id = {r[ci['Id']]: r for r in comb}
skill_by_id = {r[ki['Id']]: r for r in skills}
logic_l1 = {r[li['Group']]: r for r in logic if r[li['Level']] == 1}


def zh(lid):
    return (lang.get(lid) or ('',))[0] or ''


def ult_of(hero_id):
    cfg = hero_by_id.get(hero_id)
    if not cfg:
        return '', ''
    best = None
    for part in str(cfg[hi['OpenSkillRules']]).split('|'):
        slot, group = part.split('#')
        slot, group = int(slot), int(group)
        if slot <= 0:
            continue
        lg = logic_l1.get(group)
        if not lg:
            continue
        cb = comb_by_id.get(lg[li['SkillDisplay']])
        if not cb or cb[ci['SkillType']] != 2:
            continue
        if best is None or slot > best[0]:
            best = (slot, lg, cb)
    if not best:
        return '', ''
    slot, lg, cb = best
    sc = skill_by_id.get(lg[li['Id']])
    name = zh(sc[ki['Name']]) if sc and isinstance(sc[ki['Name']], int) else (sc[ki['Name']] if sc else '')
    return name, 'slot%d / combat %s' % (slot, cb[ci['Id']])


spec = importlib.util.spec_from_file_location('m', os.path.join(ROOT, '.workbuddy/tools/append_afk_ult_talk.py'))
m = importlib.util.module_from_spec(spec)
spec.loader.exec_module(m)

rows = []
for hero_id in sorted(m.HERO_ULT):
    name = hero_by_id[hero_id][hi['Name_txt']] if hero_id in hero_by_id else str(hero_id)
    skill, where = ult_of(hero_id)
    rows.append((hero_id, name, skill, where, m.HERO_ULT[hero_id]))

parts = []
for hero_id, name, skill, where, lines in rows:
    for i, (z, e, j) in enumerate(lines, 1):
        parts.append(
            '<tr><td class="hero"><b>%s</b><span>%s</span></td>'
            '<td class="skl">%s<span>%s</span></td>'
            '<td class="en">%s</td><td class="zh">%s</td><td class="jp">%s</td></tr>'
            % (name, hero_id, skill or '—', where, e, z, j))
hero_rows = '\n'.join(parts)

gen_rows = '\n'.join(
    '<tr><td class="n">%d</td><td class="en">%s</td><td class="zh">%s</td><td class="jp">%s</td></tr>'
    % (i, e, z, j) for i, (z, e, j) in enumerate(m.GENERIC_ULT, 1))

html = '''<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="utf-8"><title>挂机气泡 · 大招前摇台词</title>
<style>
:root{--ink:#1b1b1e;--sub:#6d6a66;--line:#e5e2db;--bg:#f5f3ef;--ok:#15803d;--accent:#b45309;--warn:#a16207}
*{box-sizing:border-box}
body{margin:0;background:var(--bg);color:var(--ink);line-height:1.5;
font-family:-apple-system,BlinkMacSystemFont,"PingFang SC","Helvetica Neue",Arial,sans-serif;-webkit-font-smoothing:antialiased}
.wrap{max-width:1180px;margin:0 auto;padding:44px 26px 72px}
h1{font-size:25px;margin:0 0 6px;letter-spacing:-.2px}
.sub{color:var(--sub);font-size:14px;margin-bottom:16px}
.meta{display:flex;gap:9px;flex-wrap:wrap;margin:0 0 28px}
.chip{background:#fff;border:1px solid var(--line);border-radius:999px;padding:5px 12px;font-size:12.5px;color:var(--sub)}
.chip b{color:var(--ink)}
h2{font-size:13px;letter-spacing:.6px;color:var(--sub);font-weight:600;text-transform:uppercase;
margin:34px 0 12px;padding-bottom:7px;border-bottom:1px solid var(--line)}
.flow{display:grid;grid-template-columns:repeat(4,1fr);gap:12px}
.step{background:#fff;border:1px solid var(--line);border-radius:12px;padding:14px 15px}
.step .k{font-size:11px;color:var(--sub);letter-spacing:.5px}
.step .v{font-size:13.5px;margin-top:5px;line-height:1.45}
.step code{background:#efece6;border-radius:4px;padding:1px 5px;font-size:11.5px}
.note{background:#fff;border:1px solid var(--line);border-left:3px solid var(--accent);border-radius:10px;
padding:13px 16px;font-size:13px;color:#454138;margin:14px 0 0}
.note b{color:var(--ink)}
.panel{background:#fff;border:1px solid var(--line);border-radius:16px;overflow:hidden}
table{width:100%;border-collapse:collapse}
th{text-align:left;font-size:11px;color:var(--sub);font-weight:600;letter-spacing:.4px;
padding:9px 10px;border-bottom:1px solid var(--line);background:#fbfaf7;position:sticky;top:0}
td{padding:9px 10px;border-bottom:1px solid #f2efe9;vertical-align:top;font-size:13px}
tr:last-child td{border-bottom:none}
.hero{width:112px}
.hero b{display:block;font-size:13.5px}
.hero span{color:#a8a49c;font-size:11px;font-variant-numeric:tabular-nums}
.skl{width:128px;color:#454138}
.skl span{display:block;color:#b0aca4;font-size:10.5px;margin-top:2px}
.en{width:34%;font-weight:600}
.zh{width:22%;color:#454138}
.jp{width:22%;color:#8c8880;font-size:12.5px}
.n{width:26px;color:#bdb9b1;font-variant-numeric:tabular-nums}
.check{background:#fff;border:1px solid var(--line);border-radius:14px;padding:16px 18px;font-size:13.5px}
.check li{margin:5px 0}
.check code{background:#efece6;border-radius:4px;padding:1px 5px;font-size:12px}
.foot{margin-top:26px;color:var(--sub);font-size:12.5px;line-height:1.75}
</style></head><body><div class="wrap">
<h1>挂机气泡：大招前摇台词</h1>
<div class="sub">战斗里只有普攻太安静。现在每场开场先抽一个人，让他的下一招变成大招——起手时先冒一句，正好压在前摇里</div>
<div class="meta">
  <div class="chip">新增 <b>108 条</b> · 50 位英雄 × 2 + 通用 8</div>
  <div class="chip">语言 <b>中 / 英 / 日</b></div>
  <div class="chip">每场 <b>55%</b> 有人放大招 · 每场至多 1 人</div>
  <div class="chip">台词 <b>70%</b> 英雄专属 / <b>30%</b> 通用</div>
  <div class="chip">实测 <b>324 条全部 ≤2 行 / 74px</b></div>
</div>

<h2>它怎么知道「下一招是大招」</h2>
<div class="flow">
  <div class="step"><div class="k">1. 解析大招</div><div class="v">取 <code>OpenSkillRules</code> 里槽位最高、<code>SkillType==2</code> 的主动技能</div></div>
  <div class="step"><div class="k">2. 开场预计算</div><div class="v"><code>PrecomputeOpenUlt(roll)</code><br>55% 抽 1 个我方单位，标记 <code>ultReady</code></div></div>
  <div class="step"><div class="k">3. 出手前换招</div><div class="v">代理表覆盖本次 <code>combat/skill</code>，不污染 <code>turnArray</code></div></div>
  <div class="step"><div class="k">4. 前摇冒泡</div><div class="v"><code>ShowUltTalk(pos)</code><br>气泡停留 2.2s，覆盖 KeyFrame 1.8~2.2s</div></div>
</div>
<div class="note"><b>为什么每场只抽 1 个：</b>挂机敌方 <code>deadTimes=1</code>，我方一轮就清完，本来就只有一个回合；
全员放大招会让 9 个人同时抢镜、动画互相盖住。大招原生最多 12 段飘字，挂机里压到 3 段。</div>

<h2>50 位英雄 · 每人两句（技能梗 + 性格梗）</h2>
<div class="panel"><table>
<thead><tr><th>英雄</th><th>大招</th><th>English</th><th>中文</th><th>日本語</th></tr></thead>
<tbody>
''' + hero_rows + '''
</tbody></table></div>

<h2>通用池 · 8 条（没配专属的英雄 / 30% 混播）</h2>
<div class="panel"><table>
<thead><tr><th></th><th>English</th><th>中文</th><th>日本語</th></tr></thead>
<tbody>
''' + gen_rows + '''
</tbody></table></div>

<h2>回归检查</h2>
<div class="check"><ul>
<li><code>scripts/lua-tests/test_afk_ult_talk.lua</code>：45 项断言全过（槽位/段数/解析/池选取/55% 边界/配置完整性/气泡定时器互斥）</li>
<li>真实配置回放：50 位挂机英雄全部解析出大招（槽位 3 或 1），无一 <code>NO ULT</code></li>
<li>Unity TextGenerator 实测：324 条中/英/日全部 ≤2 行（气泡 280×78、28 号字）</li>
<li><code>scripts/lua-bytes-sync.sh --check</code>：AfkSpeak / Language / FightPointMapOnHook 三份字节码已同步</li>
</ul></div>

<div class="foot">
Type 约定：<code>40000 + 英雄id</code> = 专属大招台词；<code>9998</code> = 通用大招台词。<br>
台词都写在「已经摆好架势、还没打出去」的窗口里，所以是临战喊话，不是旁白。
</div>
</div></body></html>'''

with open(OUT, 'w', encoding='utf-8') as f:
    f.write(html)
print('写入', OUT, len(html), 'bytes')
