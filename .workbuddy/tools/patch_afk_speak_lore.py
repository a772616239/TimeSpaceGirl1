#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""AfkSpeak 全部台词改写：贴合世界观（九大城邦 / 荒野 / 腐化者 / 旧历新历）+ 英雄自身传记梗。

结构（AfkSpeak 表，105 条）：
  Type = 0     → 通用池，任何英雄都会说（20 条）
  Type = 英雄id → 该英雄专属（85 条 + 超能DNA 20000 / 暗影梦魔 30001）
消费点：FightPointMapOnHook.HeroTalk（跑图寻敌时随机 1~2 个位置冒泡，2.5 秒）

用法：
  python .workbuddy/tools/patch_afk_speak_lore.py --dry-run
  python .workbuddy/tools/patch_afk_speak_lore.py
"""
import argparse
import io
import json
import os
import re
import subprocess
import sys

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
XLSX = os.path.join(ROOT, 'data_execl/base_data/Language.xlsx')
AFK = os.path.join(ROOT, 'data_execl/base_data/AfkSpeak.xlsx')
LUA = os.path.join(ROOT, 'Assets/ManagedResources/~Lua/Config/Data/Language.lua')
BAK = os.path.join(ROOT, '.workbuddy/backup')
PLAN = os.path.join(BAK, 'afk_lore_plan.json')

# ---------------------------------------------------------------- 通用池（Type=0）
# 世界：九大城邦 / 大都会 / 比基尔斯 / 荒野 / 腐化者 / 旧历新历 / 联邦 / 异化之海
# 语气：跑图路上随口一句；城邦的秩序 vs 荒野的烂摊子
WORLD = {
    63010089: ('腐化者跑不完，先吃饭。', 'They never end. Food first.', '腐敗者は尽きない。飯だ。'),
    63010090: ('冲！前面还是荒野。', 'Charge! More wasteland ahead.', '突撃！また荒野だ。'),
    63010091: ('荒野味。也许是血。', 'Smells like the wild. Or blood.', '荒野の匂い。血かも。'),
    63010092: ('新历了，我还是弱。', 'New Era. Same weak me.', '新暦になっても弱いまま。'),
    63010093: ('升星材料，谁去刷？', "Who's farming my star mats?", '星素材、誰が集めるんだ。'),
    63010094: ('城邦的墙真远。', 'Those city walls feel far.', '城邦の壁は遠いな。'),
    63010095: ('荒野没有外卖。', 'No delivery in the wild.', '荒野に配達はない。'),
    63010096: ('守护是城邦的玩意。', 'Guardians? City toys.', '守護は城邦の道具だ。'),
    63010097: ('九大城邦，九个老板。', 'Nine cities. Nine bosses.', '九つの城邦、九人の支配者。'),
    63010098: ('腿要断了，真的。', 'My legs are done. Truly.', '足が限界。本当に。'),
    63010099: ('遗迹在地下，不在雾里。', 'Ruins are below. Not in fog.', '遺跡は地下だ。霧ではない。'),
    63010100: ('芯片城邦出，命我自己出。', 'City gives chips. I give life.', 'チップは城邦、命は俺。'),
    63010101: ('真货都在荒野。', 'Real goods are in the wild.', '本物は荒野にある。'),
    63010102: ('神之塔：楼梯没尽头。', 'Tower of God: endless stairs.', '神の塔：終わらない階段。'),
    63010103: ('契约小字是陷阱。', 'Camp fine print is a trap.', '契約の小さい字は罠だ。'),
    63010104: ('先驱在前面，我们跟。', 'Pioneers lead. We follow.', '先駆者が先、俺らは後。'),
    63010105: ('人事处不管荒野死活。', "HR doesn't care about the wild.", '人事部は荒野を知らん。'),
    63010106: ('腐化者身上没东西掉。', 'Corrupted drop nothing. Nothing.', '腐敗者は何も落とさない。'),
    63010107: ('这图谁画的？荒野更直。', 'Bad map. The wild is straighter.', '誰が描いた地図だ。荒野の方がまし。'),
    63010108: ('星级重置？旧历的规矩。', 'Star reset? Old Era rules.', '星リセット？旧暦の掟か。'),
}

# ---------------------------------------------------------------- 英雄专属（Type=英雄id）
HERO = {
    10001: ('同行这些，差得远。', 'Peers? Not even close.', '同業者？桁が違う。'),
    10002: ('荒野比城邦安静。', 'The wild beats any city.', '荒野は城邦より静かだ。'),
    10003: ('黑帮据点，今晚全清。', 'Gang dens? Clearing tonight.', 'ギャングの拠点、今夜全部潰す。'),
    10004: ('这台机器，我发明的。', 'That machine? I built it.', 'その機械、私の発明だ。'),
    10005: ('异化之海，我说了算。', 'The Corrupted Sea is mine.', '異化の海は俺のもんだ。'),
    10006: ('爱迪生？勉强算同行。', 'Edison? Barely a rival.', 'エジソン？ただの好敵手。'),
    10007: ('机甲没电了，等我。', "Mech's out of power. Wait.", '機甲の充電が切れた。待て。'),
    10008: ('没有城邦，只有人类。', 'No cities. Just people.', '城邦ではなく、人類だけ。'),
    10009: ('一发不够，就两发。', 'One blast short? Make it two.', '一発で足りない？二発だ。'),
    10010: ('我的龙一米五，别笑。', "Dragon's 1.5m. Don't laugh.", 'うちの竜は1.5m。笑うな。'),
    10011: ('别碰我的样品。', 'Do not touch my samples.', '私のサンプルに触るな。'),
    10012: ('体术派？他们怕这剑。', 'Brawlers? They fear this sword.', '体術派？この剣を恐れてる。'),
    10013: ('谁造的我？没人知道。', 'Who made me? No one knows.', '誰が私を作った？誰も知らぬ。'),
    10014: ('富人的钱包，借一下。', "Borrowing a rich man's wallet.", '金持ちの財布、少し借りる。'),
    10015: ('这城邦太干净了。', "This city's too clean for me.", 'この城邦は綺麗すぎる。'),
    10016: ('缺什么？我店里都有。', 'Need something? My shop has it.', '何が要る？店に全部ある。'),
    10017: ('这城邦不放我的歌。', 'This city never plays my tracks.', 'この城邦、俺の曲を流さない。'),
    10018: ('砍树？先热身。', 'Chopping wood? Warm-up first.', '木を切る？まず準備運動だ。'),
    10019: ('篮球场？先热身。', 'Court? Let me warm up first.', 'コート？まず準備運動から。'),
    10020: ('值得一战的，别跑。', 'Worth fighting? Then stay.', '戦う価値がある奴、逃げるな。'),
    10021: ('这片区域，我征了。', 'This district? Conquered.', 'この区域、征服した。'),
    10022: ('跑得快就够了。', 'Being fast is enough.', '速ければ十分だ。'),
    10023: ('忍者这行，快没了。', 'Ninjas are a dying trade.', '忍の稼業も、もう終わりだ。'),
    10024: ('赵云的马，迟早我的。', "Zhao Yun's horse? Mine soon.", '趙雲の馬、いつか頂く。'),
    10025: ('九大城邦都喝我的酒。', 'Nine cities drink my wine.', '九つの城邦が俺の酒を飲む。'),
    10026: ('这球我打飞过城墙。', 'I hit one over the wall.', '城壁の向こうまで打った。'),
    10027: ('我的鹰又在骂人。', "My eagle's trash-talking again.", 'うちの鷹がまた悪口を言う。'),
    10028: ('末日过了，还是没水。', 'Apocalypse done. Still no water.', '終末は過ぎた。水はまだだ。'),
    10029: ('面罩不摘，规矩。', 'Mask stays on. House rule.', 'マスクは外さない。掟だ。'),
    10030: ('竞技场才叫活着。', 'The arena is where I live.', '闘技場でこそ生きている。'),
    10031: ('刀不沾血就生锈。', 'A blade rusts without blood.', '血を浴びねば刀は錆びる。'),
    10032: ('找人？这活我接。', "Missing person? I'm on it.", '人探し？引き受けた。'),
    10033: ('那是龙血，别争了。', 'It was dragon blood. Drop it.', '竜の血だ。もういい。'),
    10034: ('这算犯规？认真的？', 'That was a foul? Really?', '今のが反則？本気か。'),
    10035: ('我种地以来没饥荒。', 'No famine since I started.', '俺が始めてから飢饉はない。'),
    10036: ('比基尔斯，我护着。', 'Bikirsi is under my wing.', 'ビキルスは私が守る。'),
    10037: ('娱乐业？我先来的。', 'Show business? I came first.', '娯楽業？先に始めたのは私。'),
    10038: ('拳套是保护你的。', 'The glove protects you.', 'グローブはお前を守る。'),
    10039: ('那场战争，我记着。', 'I remember that war.', 'あの戦争を覚えている。'),
    10040: ('没我的宴会不算宴会。', 'No party without me.', '私のいない宴会は宴会でない。'),
    10041: ('十二试炼？全过了。', 'Twelve trials? All cleared.', '十二の試練？全部クリアだ。'),
    10042: ('右眼能看见恶魔。', 'My right eye sees demons.', '右目に悪魔が見える。'),
    10043: ('第八城邦，名义上的。', 'City Eight? Nominally mine.', '第八城邦？名目上だけだ。'),
    10044: ('猫在不在？别问。', "Is the cat here? Don't ask.", '猫はいるのか？聞くな。'),
    10045: ('你身后有个黑影。', 'There is a shadow behind you.', '後ろに黒い影がいる。'),
    10046: ('这面膜，异能做的。', 'This mask? Made with powers.', 'このパック、異能製よ。'),
    10047: ('不用异能，这是规矩。', 'No powers. That is the trick.', '異能は使わない。それが流儀。'),
    10048: ('这城邦没有秘密。', 'No secrets in this city.', 'この城邦に秘密はない。'),
    10049: ('错一个音，我听得见。', 'One wrong note. I heard it.', '一音外した。聞こえたぞ。'),
    10050: ('剪刀也能一人高。', 'Scissors can be person-sized.', 'ハサミも人ほどに大きくなる。'),
    10051: ('骑士也得先活着。', 'Even knights want to live.', '騎士だって生きていたい。'),
    10052: ('这伤，够写一章。', 'This wound? A whole chapter.', 'この傷は一章分だ。'),
    10053: ('实验室又乱了。', 'The lab is a mess. Again.', '実験室、また散らかってる。'),
    10054: ('针头专治小孩。', 'Needles fix kids fast.', '注射針は子供に効く。'),
    10057: ('我保的东西没丢过。', 'Nothing I guard gets lost.', '俺が守るものは失われない。'),
    10058: ('荒野只送尸体回来。', 'The wild sends back bodies.', '荒野から戻るのは遺体だけ。'),
    10059: ('你也许赚，我不亏。', 'You may win. I never lose.', '君は勝つかも。私は負けない。'),
    10060: ('别死，我还没下班。', "Don't die. I'm on shift.", '死ぬな。まだ勤務中だ。'),
    10061: ('球队散了？我来带。', "Team falling apart? I'll fix it.", 'チームが崩れる？私が立て直す。'),
    10062: ('守城墙，我熟。', 'Holding walls is my job.', '城壁を守るのが俺の仕事。'),
    10063: ('遗迹的画，我修的。', 'I restored those ruins.', '遺跡の絵は私が直した。'),
    10064: ('海是人类的，不是它们的。', 'The sea is ours, not theirs.', '海は人類のものだ。'),
    10065: ('统一阵线，我谈成的。', 'I got the alliance signed.', '連合をまとめたのは私だ。'),
    10066: ('这身装备，启明星定制。', 'Custom gear from Daystar.', 'この装備、明星科技の特注。'),
    10067: ('我不演，我只打分。', 'I do not perform. I judge.', '演じない。採点する。'),
    10068: ('联邦第一位女将军。', 'First female general. Mine.', '連邦初の女将軍は私だ。'),
    10069: ('星空的事，唱给你听。', 'The stars? I sing them.', '星の話は歌にする。'),
    10070: ('算式错了都找我。', 'Math broke? They call me.', '計算で困ると皆が呼ぶ。'),
    10071: ('基金会又该拨款了。', 'Time to fund the shelters.', 'そろそろ基金を回さないと。'),
    10072: ('混沌那边太顺了。', 'Chaos is going too smoothly.', '混沌、うまく行きすぎだ。'),
    10073: ('零食？给孩子的。', 'The snacks? For the kids.', 'この菓子は子供にやる。'),
    10074: ('我醒着，别怕。', "I'm awake. Mostly.", '起きてる。たぶん。'),
    10075: ('大都会，我说了算。', 'The capital answers to me.', '大都会は私の裁量だ。'),
    10076: ('秩序，得有人定。', 'Someone has to keep order.', '秩序を保つ者が要る。'),
    10079: ('我的名字，念对了？', 'You said my name right?', '私の名、正しく読めた？'),
    10080: ('结果对，手段随便。', 'Right result. Any method.', '結果が正しければ手段は問わぬ。'),
    10081: ('别动，给你穿个环。', 'Hold still. New piercing.', '動くな。ピアスを開ける。'),
    10082: ('我的事，问不得。', "Don't ask about my past.", '私の過去は聞くな。'),
    10085: ('招不招？我有耐心。', 'Talking yet? I have time.', '話すか？時間はある。'),
    10086: ('镜子里的不是我。', 'The mirror is not me.', '鏡の中のそれは私じゃない。'),
    10087: ('通缉令贴太少了。', 'Only one bounty poster?', '手配書が一枚だけか。'),
    10088: ('腐化者？到处都是。', 'Corrupted? They are everywhere.', '腐敗者？そこら中にいる。'),
    10089: ('人偶在打，我看着。', 'The puppet fights. I watch.', '人形が戦う。私は後ろで。'),
    10090: ('这不是骗术，是炼金。', 'Not a scam. Alchemy.', '詐欺じゃない。錬金術だ。'),
    20000: ('没人可用？拿我填。', 'Short a hero? Use me.', '英雄が足りない？俺で埋めろ。'),
    30001: ('我还没有名字。', 'I still have no name.', '俺にはまだ名前がない。'),
}


def lua_escape(s):
    return (s or '').replace('\\', '\\\\').replace("'", "\\'")


def read_all():
    import openpyxl
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb['Language']
    L = {}
    for r in ws.iter_rows(min_row=5, values_only=True):
        if isinstance(r[1], int):
            L[r[1]] = (r[2] or '', r[3] or '', r[4] or '')
    wb2 = openpyxl.load_workbook(AFK, data_only=True)
    ws2 = wb2['AfkSpeak']
    rows = [(r[2], r[3]) for r in ws2.iter_rows(min_row=8, values_only=True) if r[1] is not None]
    return L, rows


def build_plan(only=None):
    L, rows = read_all()
    if only:
        rows = [r for r in rows if r[0] in only]
    generic = [lid for t, lid in rows if t == 0]
    if not only and sorted(generic) != sorted(WORLD):
        print('!! 通用池 id 与 WORLD 不一致：%s' % generic)
        return None
    plan = []
    missing = []
    for t, lid in rows:
        if t == 0:
            new = WORLD[lid]
        else:
            if t not in HERO:
                missing.append(t)
                continue
            new = HERO[t]
        old = L.get(lid)
        if not old:
            missing.append(lid)
            continue
        plan.append({'id': lid, 'type': t, 'old': list(old), 'new': list(new)})
    if missing:
        print('!! 缺配置：%s' % missing)
        return None
    return plan


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--dry-run', action='store_true')
    ap.add_argument('--only', help='只处理指定 Type，逗号分隔')
    args = ap.parse_args()
    os.makedirs(BAK, exist_ok=True)

    only = None
    if args.only:
        only = {int(x) for x in args.only.split(',')}
    plan = build_plan(only)
    if plan is None:
        return 2
    print('计划 %d 条（通用 %d / 英雄 %d）' % (
        len(plan), sum(1 for p in plan if p['type'] == 0), sum(1 for p in plan if p['type'] != 0)))

    # 长度自检
    bad = []
    for p in plan:
        for i, s in enumerate(p['new']):
            lim = 34 if s.isascii() else 18
            if len(s) > lim:
                bad.append('id %s 第%d语言超长 %d>%d: %s' % (p['id'], i, len(s), lim, s))
    for b in bad:
        print(' !', b)
    if bad:
        return 2

    lua = io.open(LUA, encoding='utf-8').read()
    # old -> new，按出现次数校验
    pairs = {}
    for p in plan:
        for o, n in zip(p['old'], p['new']):
            pairs.setdefault(o, set()).add(n)
    conflicts = {o: ns for o, ns in pairs.items() if len(ns) > 1}
    if conflicts:
        for o, ns in conflicts.items():
            print(' !! 同一旧文本对应多个新文本：%s -> %s' % (o, ns))
        return 2

    replaced = 0
    for p in plan:
        for o, n in zip(p['old'], p['new']):
            lit = "'%s'" % lua_escape(o)
            cnt = lua.count(lit)
            if cnt == 0:
                print(' ! lua 找不到：%s' % o[:24])
                continue
            lua = lua.replace(lit, "'%s'" % lua_escape(n))
            replaced += cnt
    print('lua 替换字面量 %d 处' % replaced)

    json.dump([{'id': p['id'], 'old': p['old'], 'new': p['new']} for p in plan],
              open(PLAN, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)

    if args.dry_run:
        print('[DRY-RUN] 未写入')
        return 0

    io.open(LUA, 'w', encoding='utf-8', newline='').write(lua)
    print('已写入 %s' % LUA)
    r = subprocess.run([sys.executable, os.path.join(ROOT, '.workbuddy/tools/patch_language_xlsx_xml.py'),
                        '--from-json', PLAN], capture_output=True, text=True)
    print(r.stdout.strip())
    if r.returncode != 0:
        print(r.stderr.strip())
        return r.returncode
    print('下一步：scripts/lua-bytes-sync.sh Assets/ManagedResources/~Lua/Config/Data/Language.lua')
    return 0


if __name__ == '__main__':
    sys.exit(main())
