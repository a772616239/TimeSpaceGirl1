import openpyxl

file_path = '/Users/wangxufeng/Documents/TimeSpaceGirl1/data_execl/base_data/HeroRankupGroup.xlsx'
wb = openpyxl.load_workbook(file_path)
sheet = wb.active

headers = [cell.value for cell in sheet[1]]

issame_col = headers.index('Issame') + 1
issameclan_col = headers.index('IsSameClan') + 1
isid_col = headers.index('IsId') + 1

changes = 0
for row in range(8, sheet.max_row + 1):
    issame_val = sheet.cell(row=row, column=issame_col).value
    isid_val = sheet.cell(row=row, column=isid_col).value
    
    try:
        isid_int = int(isid_val) if isid_val is not None else 0
    except (ValueError, TypeError):
        isid_int = 0
        
    try:
        issame_int = int(issame_val) if issame_val is not None else 0
    except (ValueError, TypeError):
        issame_int = 0
        
    if isid_int > 0 and issame_int != 1:
        # Update values: clear specific ID requirement and add same-faction requirement
        sheet.cell(row=row, column=isid_col).value = 0
        sheet.cell(row=row, column=issameclan_col).value = 1
        changes += 1

print(f"Made {changes} changes to {file_path}.")
wb.save(file_path)

