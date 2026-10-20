# -*- coding: utf-8 -*-
"""把挂机奖励改动**精确回滚**到改动前状态。

为什么不用备份目录：`data_execl/base_data/_backup_before_tune/` 在 2026-10-20 01:30 左右被
外部删除了（原因未查明），而这 3 个 xlsx 不在 git 跟踪范围内，所以没有第二份拷贝。

好在服务端导出的 txt（`TimeSpaceServer/skzc/home/java/gameserver/conf/server/`，10-13 那次导表）
仍是**改动前**的数据，可以当作还原源：

  * `MainLevelConfig.txt` 的 RandomReward 列  -> 逐关写回 xlsx（2540 关）
  * `RewardGroup.xlsx` 删掉追加在末尾的 20101/20102/20103 三行（原组 1 与旧装备组本来就没删）
  * `RewardItemConfig.xlsx` 删掉追加在末尾的 990000~990102 共 103 行

用法：
    python scripts/adventure_reward_revert.py            # dry-run，只报告要改什么
    python scripts/adventure_reward_revert.py --apply    # 真正写回
"""
from __future__ import annotations

import argparse
import os
import shutil
import time

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BASE = os.path.join(ROOT, "data_execl", "base_data")
SERVER = os.path.join(ROOT, "TimeSpaceServer", "skzc", "home", "java", "gameserver", "conf", "server")

HEADER_ROWS = 6
NEW_GROUPS = {20101, 20102, 20103}
NEW_ITEM_IDS = set(range(990000, 990103))
SAFE_BAK = os.path.join(BASE, "_backup_before_revert")


def _num(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


def read_pristine_levels():
    """从服务端 txt 读改动前的 RandomReward。"""
    path = os.path.join(SERVER, "MainLevelConfig.txt")
    with open(path, encoding="utf-8", errors="replace") as f:
        lines = [ln.rstrip("\n").rstrip("\r") for ln in f if ln.strip()]
    header = lines[0].split("\t")
    c_id, c_rr = header.index("Id"), header.index("RandomReward")
    out = {}
    for ln in lines[2:]:
        cells = ln.split("\t")
        if len(cells) <= max(c_id, c_rr):
            continue
        lid = _num(cells[c_id])
        if lid is None:
            continue
        out[int(lid)] = cells[c_rr].replace(" ", "")
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true", help="真正写回（默认 dry-run）")
    args = ap.parse_args()

    pristine = read_pristine_levels()
    if not pristine:
        raise SystemExit("服务端 MainLevelConfig.txt 读不到数据，放弃")

    # 安全检查：如果 txt 里已经出现新组，说明它已经按新表导出过，不能再当还原源
    polluted = [lid for lid, s in pristine.items()
                if any(str(g) in s.split("#") for g in NEW_GROUPS)]
    if polluted:
        raise SystemExit(f"服务端 txt 已含新组（{polluted[:5]}...），不能作为还原源")

    f_main = os.path.join(BASE, "MainLevelConfig.xlsx")
    f_group = os.path.join(BASE, "RewardGroup.xlsx")
    f_item = os.path.join(BASE, "RewardItemConfig.xlsx")

    # ---------- MainLevelConfig ----------
    wb_main = openpyxl.load_workbook(f_main)
    ws_main = wb_main.worksheets[0]
    hdr = [c.value for c in ws_main[1]]
    c_id, c_rr = hdr.index("Id") + 1, hdr.index("RandomReward") + 1
    changed, missing = [], []
    for rw in range(HEADER_ROWS + 1, ws_main.max_row + 1):
        lid = _num(ws_main.cell(rw, c_id).value)
        if lid is None:
            continue
        lid = int(lid)
        old = pristine.get(lid)
        if old is None:
            missing.append(lid)
            continue
        cur = str(ws_main.cell(rw, c_rr).value or "").replace(" ", "")
        if cur != old:
            changed.append(lid)
            ws_main.cell(rw, c_rr).value = old

    # ---------- RewardGroup ----------
    wb_group = openpyxl.load_workbook(f_group)
    ws_group = wb_group.worksheets[0]
    group_rows = [rw for rw in range(HEADER_ROWS + 1, ws_group.max_row + 1)
                  if _num(ws_group.cell(rw, 2).value) in NEW_GROUPS]
    # 只允许删末尾连续行，避免像上次那样误删夹在中间的组
    tail_ok = group_rows and group_rows == list(range(ws_group.max_row - len(group_rows) + 1,
                                                      ws_group.max_row + 1))

    # ---------- RewardItemConfig ----------
    wb_item = openpyxl.load_workbook(f_item)
    ws_item = wb_item.worksheets[0]
    item_rows = [rw for rw in range(HEADER_ROWS + 1, ws_item.max_row + 1)
                 if _num(ws_item.cell(rw, 2).value) in NEW_ITEM_IDS]
    item_tail_ok = item_rows and item_rows == list(range(ws_item.max_row - len(item_rows) + 1,
                                                        ws_item.max_row + 1))

    print(f"MainLevelConfig：{len(changed)} 关 RandomReward 需还原"
          + (f"，{len(missing)} 关在服务端表里找不到" if missing else ""))
    print(f"RewardGroup：待删 {len(group_rows)} 行，位于末尾={bool(tail_ok)}")
    print(f"RewardItemConfig：待删 {len(item_rows)} 行，位于末尾={bool(item_tail_ok)}")

    if not args.apply:
        print("\ndry-run，未写盘。确认无误后加 --apply。")
        return 0

    if not tail_ok or not item_tail_ok:
        raise SystemExit("新增行不在表尾，脚本拒绝删除（避免误删原有数据）")

    os.makedirs(SAFE_BAK, exist_ok=True)
    stamp = time.strftime("%Y%m%d-%H%M%S")
    for f in (f_main, f_group, f_item):
        shutil.copy2(f, os.path.join(SAFE_BAK, f"{os.path.basename(f)}.{stamp}"))

    for rw in sorted(group_rows, reverse=True):
        ws_group.delete_rows(rw)
    for rw in sorted(item_rows, reverse=True):
        ws_item.delete_rows(rw)

    wb_main.save(f_main)
    wb_group.save(f_group)
    wb_item.save(f_item)

    print(f"\n已回滚。回滚前的表存到 {SAFE_BAK}")
    print("注意：客户端 Lua / 服务端 txt 仍是旧数据，无需再导表。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
