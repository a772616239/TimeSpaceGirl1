# -*- coding: utf-8 -*-
"""把新增的 RewardItemConfig 条目追加到服务端 txt。

为什么需要单独做：RewardItemConfig 是**纯后端表**（第 3 行导出类型全是 4=Server），
Unity 的 DataConfig 工具在 DataHelper.GetSheet 里遇到纯后端表会直接 return 空表，
所以「一键导表」只会写客户端 Lua 和另外两张表的 txt，这张表的 txt 不会更新。
本脚本按 DataToTxtExporter 的列顺序与格式约定把缺的行补上：

    Id \\t ItemId \\t Chance \\t GrowChance \\t RandomMin \\t RandomMax \\t Addition \\t ChanceAddition

格式约定（与现有 21292 行一致）：float 写成 x.0；空值写 null。

用法：
    python scripts/append_reward_items_to_server.py            # dry-run
    python scripts/append_reward_items_to_server.py --apply
"""
from __future__ import annotations

import argparse
import os
import shutil
import time

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BASE = os.path.join(ROOT, "data_execl", "base_data")
SERVER = os.path.join(ROOT, "TimeSpaceServer", "skzc", "home", "java", "gameserver", "conf", "server")
SERVER_DEPLOY = os.path.join(ROOT, "TimeSpaceServer", "Server_x1tag1.2.2", "server_game", "conf", "server")

HEADER_ROWS = 6
NEW_IDS = range(990000, 990103)


def _num(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


def fmt_float(v):
    return f"{v:.1f}"


def build_rows():
    ws = openpyxl.load_workbook(os.path.join(BASE, "RewardItemConfig.xlsx")).worksheets[0]
    out = []
    for r in range(HEADER_ROWS + 1, ws.max_row + 1):
        rid = _num(ws.cell(r, 2).value)
        if rid is None or int(rid) not in NEW_IDS:
            continue
        item_id = int(_num(ws.cell(r, 3).value) or 0)
        chance = int(_num(ws.cell(r, 4).value) or 0)
        grow = _num(ws.cell(r, 5).value)
        grow = 1.0 if grow is None else grow
        lo = int(_num(ws.cell(r, 6).value) or 0)
        hi = int(_num(ws.cell(r, 7).value) or 0)
        add = ws.cell(r, 8).value
        cadd = ws.cell(r, 9).value
        add = "null" if add is None or str(add).strip() == "" else str(add).strip()
        cadd = "null" if cadd is None or str(cadd).strip() == "" else str(cadd).strip()
        out.append((int(rid), f"{int(rid)}\t{item_id}\t{chance}\t{fmt_float(grow)}\t{lo}\t{hi}\t{add}\t{cadd}"))
    return out


def existing_ids(path):
    ids = set()
    with open(path, encoding="utf-8", errors="replace") as f:
        for i, ln in enumerate(f):
            if i < 2 or not ln.strip():
                continue
            v = ln.split("\t")[0]
            if v.isdigit():
                ids.add(int(v))
    return ids


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--sync-deploy", action="store_true",
                    help="同时把三张表复制到 Server_x1tag1.2.2/server_game/conf/server")
    args = ap.parse_args()

    rows = build_rows()
    if len(rows) != len(NEW_IDS):
        raise SystemExit(f"xlsx 里只找到 {len(rows)} 条新增条目，预期 {len(NEW_IDS)}，放弃")

    path = os.path.join(SERVER, "RewardItemConfig.txt")
    have = existing_ids(path)
    todo = [(rid, line) for rid, line in rows if rid not in have]
    print(f"服务端已有 {len(have & set(NEW_IDS))} 条新条目，待追加 {len(todo)} 条")
    if todo:
        print("示例：")
        for _, line in todo[:3]:
            print("   " + line)

    if not args.apply:
        print("\ndry-run，未写盘。")
        return 0

    if todo:
        with open(path, "rb") as f:
            f.seek(-1, os.SEEK_END)
            tail = f.read(1)
        with open(path, "a", encoding="utf-8") as f:
            if tail not in (b"\n", b"\r"):
                f.write("\n")
            for _, line in todo:
                f.write(line + "\n")
        print(f"已追加 {len(todo)} 行到 {os.path.relpath(path, ROOT)}")

    # 回读校验
    ids = existing_ids(path)
    missing = [i for i in NEW_IDS if i not in ids]
    if missing:
        raise SystemExit(f"校验失败，仍缺 {len(missing)} 条：{missing[:5]}")
    print(f"校验通过：{len(NEW_IDS)} 条新条目全部在服务端表中")

    if args.sync_deploy:
        stamp = time.strftime("%Y%m%d-%H%M%S")
        bak = os.path.join(SERVER_DEPLOY, f"_bak_{stamp}")
        os.makedirs(bak, exist_ok=True)
        for name in ("MainLevelConfig.txt", "RewardGroup.txt", "RewardItemConfig.txt"):
            src = os.path.join(SERVER, name)
            dst = os.path.join(SERVER_DEPLOY, name)
            if os.path.exists(dst):
                shutil.copy2(dst, os.path.join(bak, name))
            shutil.copy2(src, dst)
            print(f"同步 {name} -> server_game（旧文件备份在 {os.path.relpath(bak, ROOT)}）")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
