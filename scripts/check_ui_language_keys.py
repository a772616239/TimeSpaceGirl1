#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""巡检：UI 预制体里还没进语言表的中文文本。

背景（2026-10-22）：
    UIManager.MultiLanguageCheck 在面板打开时，对每个 Text **按中文原文**去
    languageDicStr 查表（键 = Language.xlsx 的 Chinese 列）。所以「预制体里写了
    中文、语言表里没有这一行」⇒ 英文/日文下原样显示中文，控制台只留一行
    `languageDicStr no find Str: xxx`，肉眼极难发现。

用法：
    python scripts/check_ui_language_keys.py                # 全量巡检（默认 Prefabs/UI + UpdatePanel）
    python scripts/check_ui_language_keys.py --path Assets/ManagedResources/Prefabs/UI/Lead
    python scripts/check_ui_language_keys.py --fail         # 有缺失时返回 1（给 CI 用）
    python scripts/check_ui_language_keys.py --csv out.csv  # 落表，方便逐条翻译

注意：
    - 只统计**含中文**的 m_Text；数字/纯符号/富文本标签不参与。
    - 文本里 \\uXXXX 转义要先解码，否则永远匹配不上。
    - 「预制体写了中文但运行时会被代码覆盖」的占位文本也会被报出来，属预期噪音，
      逐条判断，别盲目全加语言表。
"""
import argparse
import csv
import os
import re
import sys

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
XLSX = os.path.join(ROOT, 'data_execl/base_data/Language.xlsx')
DEFAULT_ROOTS = [
    'Assets/ManagedResources/Prefabs/UI',
    'Assets/ManagedResources/UpdatePanel',
]
CJK = re.compile(r'[\u4e00-\u9fa5]')
TEXT_ESC = re.compile(r'\\u([0-9a-fA-F]{4})')


def load_zh_keys(path=XLSX):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    keys = set()
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[1] is None:
            continue
        zh = row[2]
        if isinstance(zh, str) and zh:
            keys.add(zh)
    wb.close()
    return keys


def decode(text):
    return TEXT_ESC.sub(lambda m: chr(int(m.group(1), 16)), text)


def scan_prefab(path):
    """返回 [(text, component_name)]"""
    out = []
    with open(path, encoding='utf-8', errors='replace') as f:
        raw = f.read()
    for block in re.split(r'\n--- !u!', raw):
        m = re.search(r'^\s*m_Text: (.*)$', block, re.M)
        if not m:
            continue
        val = m.group(1).strip()
        if len(val) >= 2 and val[0] == '"' and val[-1] == '"':
            val = val[1:-1]
        else:
            continue
        val = decode(val)
        if not CJK.search(val):
            continue
        out.append((val, ''))
    return out


def walk(roots):
    for root in roots:
        full = root if os.path.isabs(root) else os.path.join(ROOT, root)
        if os.path.isfile(full):
            yield full
            continue
        for dirpath, _dirs, files in os.walk(full):
            for name in sorted(files):
                if name.endswith('.prefab'):
                    yield os.path.join(dirpath, name)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--path', action='append', default=None,
                    help='只扫指定目录/文件（可重复）')
    ap.add_argument('--fail', action='store_true', help='有缺失时返回 1')
    ap.add_argument('--csv', help='把缺失清单写成 csv')
    ap.add_argument('--limit', type=int, default=0, help='每个预制体最多打印几条')
    args = ap.parse_args()

    roots = args.path or DEFAULT_ROOTS
    keys = load_zh_keys()
    missing = {}
    total_text = 0

    for prefab in walk(roots):
        hits = []
        for text, _ in scan_prefab(prefab):
            total_text += 1
            if text not in keys:
                hits.append(text)
        if hits:
            missing[os.path.relpath(prefab, ROOT)] = hits

    files = len(missing)
    texts = sum(len(v) for v in missing.values())
    uniq = sorted({t for v in missing.values() for t in v})
    print('语言表中文键：%d 条；扫描含中文文本 %d 处' % (len(keys), total_text))
    print('未进语言表：%d 个预制体 / %d 处文本 / %d 条唯一文案'
          % (files, texts, len(uniq)))
    for path in sorted(missing):
        print('\n== %s' % path)
        items = missing[path]
        show = items if not args.limit else items[:args.limit]
        for text in show:
            print('   - %s' % text.replace('\n', '\\n'))
        if args.limit and len(items) > args.limit:
            print('   … 另有 %d 处' % (len(items) - args.limit))

    if args.csv:
        with open(args.csv, 'w', newline='', encoding='utf-8-sig') as f:
            w = csv.writer(f)
            w.writerow(['prefab', 'zh', 'en', 'jp'])
            for path in sorted(missing):
                for text in missing[path]:
                    w.writerow([path, text, '', ''])
        print('\n已写出：%s' % args.csv)

    if args.fail and files:
        return 1
    return 0


if __name__ == '__main__':
    sys.exit(main())
