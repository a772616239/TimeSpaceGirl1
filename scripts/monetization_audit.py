#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""充值全表审计（只读）。

把「充值相关表」的体检项固化成可复跑脚本，改动前/改动后各跑一次即可对比。

用法：
    python3 scripts/monetization_audit.py            # 全量审计
    python3 scripts/monetization_audit.py --section value   # 只看价值锚点
    python3 scripts/monetization_audit.py --json out.json   # 机器可读

审计项：
    A 货架       SKU 总数 / 可售 / 门控死行 / 整条关闭的产品线
    B 计费点     每个渠道的覆盖数（0 覆盖 = 该渠道无法充值）
    C 价值锚点   首充 / 月卡赠品 / 基金 / 战令 / 礼包带 / 直充 的抽等值每元
    D 倒挂       同族大额档低于小额档（系统性定价缺陷）
    E 同价离散   同一价位内不同礼包的抽等值离散度
    F 钻石面     商店券价 vs 直抽成本 / 招财猫 ROI / VIP 日赠
    G 空包       在售但价值字段为空的 SKU

口径：
    抽等值 = 券(19/20/87/50007) 张数 + 钻石(16) / 220
    220 = 单抽直耗钻；十连 2000 钻 = 200/抽。
    不计工坊币、经验、金币等非抽卡资源——本表只回答「同样的钱能换到几抽」。
"""

from __future__ import annotations

import argparse
import collections
import json
import os
import statistics
import sys

import openpyxl

BASE = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                    "data_execl", "base_data")

DIAMOND = "16"
DP_PER_PULL = 220          # 单抽直耗钻石
DP_TEN_PULL = 2000         # 十连直耗钻石
TICKETS = ("19", "20", "87", "50007")
CURRENCY_NAME = {"16": "钻石", "19": "特殊招募券", "20": "普通招募券",
                 "87": "券87", "50007": "券50007", "14": "金币"}

# 首充 / 成长基金 / 基金 / 战令 的锚点 SKU
ANCHOR_SKUS = {1: "首充 6 元", 101: "成长基金 98 元",
               78: "豪华基金 128 元", 79: "豪华基金 328 元",
               72001: "战令 198 元"}


# ─────────────────────────── 基础工具 ───────────────────────────

def load(name: str, sheet: str | None = None):
    path = os.path.join(BASE, name + ".xlsx")
    if not os.path.exists(path):
        sys.exit("缺少配置表：%s" % path)
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.worksheets[0]
    return wb, ws


def header_index(ws) -> dict:
    hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    return {h: i for i, h in enumerate(hdr) if h}


def iter_data(ws, first_row: int = 8):
    for r in ws.iter_rows(min_row=first_row, values_only=True):
        yield r


def pairs(text) -> collections.Counter:
    """"19#10|16#600" -> Counter({"19":10,"16":600})"""
    out = collections.Counter()
    for chunk in str(text or "").split("|"):
        if "#" not in chunk:
            continue
        key, _, num = chunk.partition("#")
        try:
            out[key.strip()] += int(float(num))
        except (TypeError, ValueError):
            continue
    return out


def pull_value(counter: collections.Counter) -> float:
    tickets = sum(counter.get(k, 0) for k in TICKETS)
    return tickets + counter.get(DIAMOND, 0) / DP_PER_PULL


def rule_gated(rule) -> bool:
    """OnSaleRule = 1#999#999 表示等级 999 才能看到，等价于永久关停。"""
    return str(rule or "").replace(" ", "").startswith("1#999#999")


# ─────────────────────────── 读表 ───────────────────────────

def load_skus():
    """RechargeCommodityConfig 全量 + 计费点覆盖标记。"""
    wb, ws = load("ChannelRechargeIdConfig")
    hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    channels = [h for h in hdr[3:] if h]
    billing, channel_cov = {}, collections.Counter()
    for r in iter_data(ws):
        sid = r[1]
        if sid is None:
            continue
        flags = [x not in (None, "", 0) for x in r[3:3 + len(channels)]]
        billing[sid] = any(flags)
        for k, ok in enumerate(flags):
            if ok:
                channel_cov[channels[k]] += 1
    wb.close()

    wb2, ws2 = load("RechargeCommodityConfig")
    idx = header_index(ws2)
    skus = {}
    for r in iter_data(ws2, first_row=6):
        sid = r[idx["Id"]]
        if sid is None:
            continue
        base = pairs(r[idx["BaseReward"]])
        ext = pairs(r[idx["ExtraReward"]])
        first = pairs(r[idx["FirstMultiple"]])
        total = base + ext
        skus[sid] = {
            "id": sid,
            "name": r[idx["Name"]],
            "price": r[idx["Price"]] or 0,
            "type": r[idx["Type"]],
            "show_type": r[idx["ShowType"]],
            "limit": r[idx["Limit"]],
            "time": r[idx["Time"]],
            "otype": r[idx["Otype"]],
            "rule": r[idx["OnSaleRule"]],
            "base": base,
            "extra": ext,
            "first": first,
            "pull": pull_value(total),
            "pull_first": pull_value(total + first),
            "billing": billing.get(sid, False),
            "gated": rule_gated(r[idx["OnSaleRule"]]),
        }
    wb2.close()
    return skus, channels, channel_cov


def load_fund_totals():
    """LuxuryFund 按活动类型汇总 30 天总产出。"""
    wb, ws = load("LuxuryFund")
    out = collections.defaultdict(collections.Counter)
    for r in iter_data(ws):
        if r[1] is None:
            continue
        out[r[2]] += pairs(r[4])
    wb.close()
    return out


def load_activity_tracks():
    """活动奖励轨道：activityId -> 抽等值总量，以及 skuId -> activityId 的反查表。

    成长基金（Type=4）的 SKU 行是空的，价值全在 ActivityRewardConfig 的节点里，
    由 GlobalActivity.CanBuyRechargeId 挂到 SKU 上。漏掉这一步会把基金误判成空包。
    """
    wb, ws = load("ActivityRewardConfig")
    idx = header_index(ws)
    track = collections.Counter()
    for r in iter_data(ws):
        if r[idx["Id"]] is None:
            continue
        track[r[idx["ActivityId"]]] += pull_value(pairs(r[idx["Reward"]]))
    wb.close()

    wb2, ws2 = load("GlobalActivity")
    idx2 = header_index(ws2)
    sku_to_act = {}
    for r in iter_data(ws2):
        if r[idx2["Id"]] is None:
            continue
        buy = str(r[idx2.get("CanBuyRechargeId")] or "").strip()
        if buy.isdigit():
            sku_to_act[int(buy)] = r[idx2["Id"]]
    wb2.close()
    return track, sku_to_act


def sku_total_pull(sku, funds, tracks, sku_to_act):
    """一个 SKU 的完整抽等值：SKU 行 + 基金日领 + 活动轨道。"""
    total = sku["pull"]
    if sku["id"] in (78, 79) and sku["id"] in funds:
        total += pull_value(funds[sku["id"]])
    act = sku_to_act.get(sku["id"])
    if act and act in tracks:
        total += tracks[act]
    return total


def live_skus(skus):
    """可售 = 有计费点 + 未门控 + 价格大于 0。"""
    return [s for s in skus.values()
            if s["billing"] and not s["gated"] and s["price"] > 0]


# ─────────────────────────── 审计项 ───────────────────────────

def section_shelf(skus):
    live = live_skus(skus)
    gated = [s for s in skus.values() if s["gated"]]
    print("A. 货架清点")
    print("   SKU 总行数            %d" % len(skus))
    print("   可售（有能力被买）    %d  (%.1f%%)" % (len(live), len(live) * 100.0 / len(skus)))
    print("   999 门控死行          %d  (%.1f%%)" % (len(gated), len(gated) * 100.0 / len(skus)))
    print("     其中同时无计费点    %d" % sum(1 for s in gated if not s["billing"]))

    live_show, dead_show = collections.Counter(), collections.Counter()
    for s in skus.values():
        alive = s["billing"] and not s["gated"] and s["price"] > 0
        (live_show if alive else dead_show)[s["show_type"]] += 1
    shut = sorted(k for k in dead_show if live_show.get(k, 0) == 0)
    print("   整条关闭的 ShowType   %s" % (shut if shut else "无"))
    print("   → 这些产品线在配置里仍有价格和内容，只是被门控关停；"
          "若属于「待上线」应登记，若属于「已废弃」应从表里清掉，否则会持续误导数值对比。")
    return live, gated, shut


def section_billing(skus, channels, channel_cov):
    print("\nB. 计费点覆盖（渠道侧能不能真的付钱）")
    total = len(skus)
    dead = [ch for ch in channels if channel_cov[ch] == 0]
    for ch in channels:
        n = channel_cov[ch]
        flag = "  ◀ 全渠道 0 覆盖，该渠道无法充值" if n == 0 else ""
        print("   %-30s %4d / %d%s" % (ch, n, total, flag))
    return dead


def section_value(skus, funds, tracks, sku_to_act):
    print("\nC. 价值锚点（抽等值 / 元）")
    print("   %-24s %-8s %-10s %-9s %s" % ("产品", "价格", "抽等值", "抽/元", "备注"))
    rows = []

    for sid, label in ANCHOR_SKUS.items():
        s = skus.get(sid)
        if not s:
            continue
        total_pull = sku_total_pull(s, funds, tracks, sku_to_act)
        if sid == 1:                   # 首充，含 FirstMultiple 加成
            total_pull = s["pull_first"]
        if sid == 72001:               # 战令 = 购买礼包 + 任务轨道
            wb, ws = load("BattlePassReward")
            track = collections.Counter()
            for r in iter_data(ws):
                if r[1] is not None:
                    track += pairs(r[3])
            wb.close()
            total_pull += pull_value(track)
        rate = total_pull / s["price"] if s["price"] else 0
        rows.append((label, s["price"], total_pull, rate))
        print("   %-24s %-8s %-10.1f %-9.4f" % (label, s["price"], total_pull, rate))

    # 月卡：SKU 被门控，真实形态是 MonthcardConfig 的累充赠品
    wb, ws = load("MonthcardConfig")
    for r in iter_data(ws):
        if r[1] is None:
            continue
        gate, per_day, days = r[4], pairs(r[5]), r[6] or 0
        if not gate or not days:
            continue
        total = pull_value(per_day) * days
        print("   %-22s %-9s %-11.1f %-9.4f %s"
              % ("%s（累充赠送）" % r[2], gate, total, total / gate,
                 "门槛 %d 元，白拿 %.1f 抽/月" % (gate, total)))
    wb.close()

    # 礼包带中位 & 直充
    pack = [s for s in live_skus(skus) if s["type"] == 5]
    band = collections.defaultdict(list)
    for s in pack:
        band[s["price"]].append(s["pull"] / s["price"])
    print("   ── 可复购礼包带（Type=5 中位）")
    for p in sorted(band):
        print("   %-22s %-9s %-11s %-9.4f" % ("礼包 %d 元" % p, p, "%d 个" % len(band[p]),
                                             statistics.median(band[p])))
    print("   ── 直充（Type=3，无限次）")
    for s in sorted((s for s in live_skus(skus) if s["type"] == 3), key=lambda x: x["price"]):
        print("   %-22s %-9s %-11.1f %-9.4f  首购 %.4f"
              % ("直充 %d 元" % s["price"], s["price"], s["pull"],
                 s["pull"] / s["price"], s["pull_first"] / s["price"]))

    # 阶梯落差：最好的一档比最差的一档高多少倍
    rated = [(lab, p, r) for lab, p, _v, r in rows if p]
    if rated and band:
        top = max(rated, key=lambda x: x[2])
        pack_mid = statistics.median([s["pull"] / s["price"]
                                      for s in live_skus(skus) if s["type"] == 5])
        print("   ── 落差：最优锚点 %s %.4f 抽/元 = 礼包带中位 %.4f 的 %.1f 倍"
              % (top[0], top[2], pack_mid, top[2] / pack_mid))
        print("   → 一次档/进度档比可复购带高 4 倍以上时，可复购货架在数值上就没有购买理由。")
    return rows


def section_inversion(live):
    print("\nD. 同族倒挂（同一 ShowType 内，大额档的抽/元低于小额档）")
    fam = collections.defaultdict(list)
    for s in live:
        fam[s["show_type"]].append(s)
    hits = 0
    for st, rows in sorted(fam.items()):
        if len(rows) < 2:
            continue
        by_price = collections.defaultdict(list)
        for s in rows:
            by_price[s["price"]].append(s)
        prices = sorted(by_price)
        med = {p: statistics.median([x["pull"] / x["price"] for x in by_price[p]])
               for p in prices}
        for a, b in zip(prices, prices[1:]):
            if med[b] < med[a] - 1e-9:
                hits += 1
                print("   ✗ ShowType=%-4s %s 元 %.4f → %s 元 %.4f  (跌 %.1f%%)"
                      % (st, a, med[a], b, med[b], (1 - med[b] / med[a]) * 100))
    if not hits:
        print("   无")
    print("   合计 %d 处" % hits)
    return hits


def section_dispersion(live):
    print("\nE. 同价离散（同一价位内，不同礼包的价值差多少倍）")
    by_price = collections.defaultdict(list)
    for s in live:
        if s["type"] == 5:
            by_price[s["price"]].append(s)
    worst = []
    for p in sorted(by_price):
        rates = [s["pull"] / s["price"] for s in by_price[p]]
        if len(rates) < 3:
            continue
        lo, hi = min(rates), max(rates)
        if lo <= 0:
            continue
        spread = hi / lo
        worst.append((spread, p, lo, hi, len(rates)))
        if spread >= 2.0:
            print("   ✗ %-5s 元 %-3d 个：%.4f ~ %.4f  = %.1fx" % (p, len(rates), lo, hi, spread))
    if worst:
        worst.sort(reverse=True)
        top = worst[0]
        print("   最严重：%s 元档内高低差 %.1f 倍（%.4f vs %.4f）" % (top[1], top[0], top[2], top[3]))
    print("   → 同一价位里价值差 2 倍以上，等于让玩家抽奖决定买哪个；"
          "通常应把同档收敛到中位 ±20%。")
    return worst


def section_diamond(live, skus):
    print("\nF. 钻石面（回收价与产出）")

    # F1 商店券价 vs 直抽
    wb, ws = load("StoreConfig")
    idx = header_index(ws)
    print("   ── 商店卖券的钻石单价（对照直抽 %d 钻/抽、十连 %d 钻/抽）"
          % (DP_PER_PULL, DP_TEN_PULL // 10))
    seen = set()
    for r in iter_data(ws):
        if r[idx["Id"]] is None:
            continue
        goods = str(r[idx["Goods"]] or "")
        cost = str(r[idx["Cost"]] or "")
        if "|" in goods:            # 组合包，单价没有比较意义
            continue
        item = goods.split("#")[0]
        if item not in ("19", "20") or not cost.startswith(DIAMOND + "|"):
            continue
        try:
            qty = int(float(goods.split("#")[1]))
            price = float(cost.split("|")[1].split("#")[-1])
        except (IndexError, ValueError):
            continue
        if qty <= 0:
            continue
        unit = price / qty
        key = (item, unit)
        if key in seen:
            continue
        seen.add(key)
        if item == "19" and unit > DP_PER_PULL:
            print("   ✗ 券19 单价 %.0f 钻/张 > 直抽 %d 钻/抽（买券即亏 %.1f%%）"
                  % (unit, DP_PER_PULL, (unit / DP_PER_PULL - 1) * 100))
        elif item == "19" and unit < DP_TEN_PULL // 10:
            print("   ✗ 券19 单价 %.0f 钻/张 < 十连 %d 钻/抽（存在套利）"
                  % (unit, DP_TEN_PULL // 10))
        else:
            print("   · 券%-3s 单价 %-6.0f 钻/张" % (item, unit))
    wb.close()

    # F2 招财猫
    wb, ws = load("LuckyCatConfig")
    agg = collections.defaultdict(lambda: [0, 0])
    for r in iter_data(ws):
        if r[1] is None or not r[5]:
            continue
        agg[(r[2], r[3])][0] += r[5]
        agg[(r[2], r[3])][1] += r[6]
    wb.close()
    print("   ── 招财猫（消耗钻石换钻石）")
    for (act, _typ), (cost, up) in sorted(agg.items()):
        if cost:
            print("   活动 %-6s 满次数消耗 %-6d → 上限收益 %-6d = %.2fx%s"
                  % (act, cost, up, up / cost, "  ◀ 净产出钻石" if up > cost else ""))
    print("   → ROI > 1 的「回收口」实际是产出通道，会把任何付费钻石按倍率放大。")

    # F3 VIP 日赠
    wb, ws = load("VipLevelConfig")
    daily = []
    for r in iter_data(ws):
        if r[1] is None or not r[3]:
            continue
        daily.append((r[1], pull_value(pairs(r[3]))))
    wb.close()
    if daily:
        print("   ── VIP 每日礼包（永久日赠）")
        print("   VIP%-3s %.2f 抽/天   →   VIP%-3s %.2f 抽/天   （%.0f 倍差）"
              % (daily[0][0], daily[0][1], daily[-1][0], daily[-1][1],
                 daily[-1][1] / max(daily[0][1], 1e-9)))


def section_empty(live, funds, tracks, sku_to_act):
    print("\nG. 空包（在售，但 SKU 行 / 基金日领 / 活动轨道三处都拿不到价值）")
    hits = 0
    for s in live:
        if sku_total_pull(s, funds, tracks, sku_to_act) > 0 or s["base"] or s["extra"]:
            continue
        hits += 1
        print("   ✗ id=%-6s %s 元  Type=%s ShowType=%s limit=%s  base/extra/轨道 全空"
              % (s["id"], s["price"], s["type"], s["show_type"], s["limit"]))
    if not hits:
        print("   无")
    else:
        print("   → 价格字段有值但三处都无奖励，付款后没有任何产出。上线前必须补齐或门控关闭。")
    return hits


# ─────────────────────────── 入口 ───────────────────────────

def main():
    ap = argparse.ArgumentParser(description="充值全表审计（只读）")
    ap.add_argument("--section", choices=["shelf", "billing", "value", "inversion",
                                          "dispersion", "diamond", "empty"],
                    help="只跑某一项")
    ap.add_argument("--json", help="把汇总结果写到 JSON")
    args = ap.parse_args()

    skus, channels, channel_cov = load_skus()
    funds = load_fund_totals()
    tracks, sku_to_act = load_activity_tracks()
    live = live_skus(skus)
    only = args.section
    summary = {}

    if only in (None, "shelf"):
        _live, gated, shut = section_shelf(skus)
        summary["shelf"] = {"total": len(skus), "live": len(live),
                            "gated": len(gated), "shut_lines": len(shut)}
    if only in (None, "billing"):
        summary["billing_zero_channels"] = section_billing(skus, channels, channel_cov)
    if only in (None, "value"):
        section_value(skus, funds, tracks, sku_to_act)
    if only in (None, "inversion"):
        summary["inversions"] = section_inversion(live)
    if only in (None, "dispersion"):
        w = section_dispersion(live)
        summary["worst_spread"] = w[0][0] if w else None
    if only in (None, "diamond"):
        section_diamond(live, skus)
    if only in (None, "empty"):
        summary["empty_packs"] = section_empty(live, funds, tracks, sku_to_act)

    if args.json:
        with open(args.json, "w", encoding="utf-8") as fh:
            json.dump(summary, fh, ensure_ascii=False, indent=2)
        print("\n汇总已写入 %s" % args.json)


if __name__ == "__main__":
    main()
