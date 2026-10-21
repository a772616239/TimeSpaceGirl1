#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""R13 结构性定价修复：把「一次性档 / 进度档」的价值落差从 5.6 倍收到 2.9 倍。

四个改动项（--only 可单选，默认全做）：

  fund    成长基金 活动 6/601/602/603/604 的 10 个节点统一 16#2156
          30,160 钻/98 元（1.399 抽/元）→ 21,560 钻/98 元（1.000 抽/元）
  pack    Type=5 直购礼包整体 ×1.21（**只加券，一颗钻石不加**）
          648 档 0.2915 → ≈0.353 抽/元
  store   商店里以钻石标价的 19#N 券，单价从 250 → 220（10 连 2500 → 2000）
          修正「买券比直接抽卡还贵」的价格倒挂
  lux328  豪华基金 328 元：LuxuryFund Type=79 的 17 天 16#2080 → 16#2540
          0.7184 → 0.8268 抽/元，与 128 元档 0.8267 持平

不变的东西（自检逐条守）：
  · `Price` 一律不动 —— 计费点按商品 id 绑定，改价会让渠道侧金额与配置不符。
  · 钻石产出不动（fund 项除外，那是它的本体）—— 钻石是全局通货。
  · 任何 SKU 的「非券道具构成」不动 —— 只往 BaseReward/RewardShow 里加券。
  · 商店只降价不涨价。

用法：
    python3 scripts/monetization_rebalance.py                  # dry-run
    python3 scripts/monetization_rebalance.py --apply          # 备份+落表+自检（失败自动回滚）
    python3 scripts/monetization_rebalance.py --verify-only
    python3 scripts/monetization_rebalance.py --only fund,store
    python3 scripts/monetization_rebalance.py --revert <backup_dir>
    python3 scripts/monetization_rebalance.py --list-backups
"""

from __future__ import annotations

import argparse
import collections
import os
import shutil
import statistics
import sys
import time

try:
    import openpyxl
except ImportError:
    sys.exit("需要 openpyxl：/Users/wangxufeng/.workbuddy/binaries/python/envs/default/bin/python")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BASE = os.path.join(ROOT, "data_execl", "base_data")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from backup_paths import sub as _bak_sub              # noqa: E402
from monetization_ladder import (                     # noqa: E402
    DIA, DIA_PER_PULL, TICKETS, add_items, fmt_pairs,
    header_index, parse_pairs, pick_ticket, pull_equiv,
)

BACKUP_ROOT = _bak_sub("_backup_before_rebalance")
# 幂等锁：本脚本的 pack 项是**乘法**，连跑两次会把礼包带乘到 1.21²。
# 每项落表成功后写一条记录，再次 --apply 同一项必须先 --force。
MARKER = os.path.join(BACKUP_ROOT, "_applied.json")

F_REWARD = "ActivityRewardConfig"
F_RECHARGE = "RechargeCommodityConfig"
F_STORE = "StoreConfig"
F_FUND = "LuxuryFund"

# 各表的数据起始行（与导表工具的约定一致：Recharge 是 6，其余是 8）
ROW_START = {F_RECHARGE: 6, F_REWARD: 8, F_STORE: 8, F_FUND: 8}

# ─────────────────────────── fund ───────────────────────────
# 成长基金：SKU id -> 活动 Id（OnSaleRule 的 `3#N` 就是活动号，verify 会核对）
FUND_SKU_ACT = {101: 6, 102: 601, 103: 602, 104: 603, 105: 604}
FUND_NODE_DIA = 2156          # 每节点钻石数；10 节点 = 21,560 = 98 抽 = 1.00 抽/元
FUND_PRICE = 98

# ─────────────────────────── pack ───────────────────────────
PACK_TYPE = 5
PACK_MULT = 1.21
# 券的最小粒度就是 1 张。低价包上补 1 张等于 +65%（6 元）/+32%（12 元），
# 会造成"6 元包比 648 元包值"的反向阶梯 —— 所以只有实际增幅落在
# [MULT-TOL, MULT+TOL] 内的 SKU 才动，其余如实标注"粒度不足"。
PACK_REL_TOL = 0.10
# 整档中位抽数低于此值就不动：补 1 张券的相对误差会超过 20%。
PACK_MIN_TIER_PULL = 5.0
# 战令：Type=5 但价值大头在 BattlePassReward 任务轨道（不受本脚本控制），
# 单独乘 1.21 会让「SKU 值」和「轨道值」脱节，故排除，由战令自己定档。
PACK_EXCLUDE = {72001}

# ─────────────────────────── store ───────────────────────────
TICKET_GOODS = "19"           # 只修紫晶卡：它是唯一有「直抽钻石价」可比的券
STORE_UNIT = 220              # 单抽直耗钻（LotterySetting Id=12：19#1|16#220）
STORE_TEN_UNIT = 200          # 十连均摊（Id=11：19#10|16#2000）

# ─────────────────────────── lux328 ───────────────────────────
LUX_TYPE_OLD = 79
LUX_DAY_OLD = 2080
LUX_DAY_NEW = 2540
LUX_SKU = 79
LUX_SKU_SHOW = "16#38640"


# ─────────────────────────── 通用 ───────────────────────────

def path_of(name):
    p = os.path.join(BASE, name + ".xlsx")
    if not os.path.exists(p):
        sys.exit("缺少配置表：%s" % p)
    return p


def load(name, sheet=None):
    wb = openpyxl.load_workbook(path_of(name))
    ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.worksheets[0]
    return wb, ws


def iter_data(ws, start):
    for r in range(start, ws.max_row + 1):
        v = [c.value for c in ws[r]]
        if len(v) < 2 or v[1] is None:
            continue
        yield r, v


def pairs_dict(text):
    return {k: (n or 0) for k, n in parse_pairs(text)}


def backup(paths):
    ts = time.strftime("%Y%m%d-%H%M%S")
    d = os.path.join(BACKUP_ROOT, ts)
    os.makedirs(d, exist_ok=True)
    for p in paths:
        shutil.copy2(p, os.path.join(d, os.path.basename(p)))
    return d


def _restore(bak):
    for f in os.listdir(bak):
        shutil.copy2(os.path.join(bak, f), os.path.join(BASE, f))
    print("   已回滚 ← %s" % bak)


def load_marker():
    import json
    if not os.path.exists(MARKER):
        return {}
    try:
        with open(MARKER, "r", encoding="utf-8") as fh:
            return json.load(fh)
    except (OSError, ValueError):
        return {}


def save_marker(items, bak_dir):
    import json
    data = load_marker()
    for it in items:
        data[it] = {"applied_at": time.strftime("%Y-%m-%dT%H:%M:%S"), "backup": bak_dir}
    os.makedirs(os.path.dirname(MARKER), exist_ok=True)
    with open(MARKER, "w", encoding="utf-8") as fh:
        json.dump(data, fh, ensure_ascii=False, indent=2)


def guard(items, force):
    """乘法项重复落表 = 重复加息。已经有记录就必须 --force。"""
    done = load_marker()
    hit = [it for it in items if it in done]
    if hit and not force:
        sys.exit("❌ %s 已经落过表了（%s）。\n"
                 "   本脚本的 pack 项是乘法，再跑一次会再乘 1.21。\n"
                 "   确认要重做请先 --revert <备份目录> 再 --apply --force。"
                 % (",".join(hit), ", ".join(done[h]["applied_at"] for h in hit)))


class Checker:
    def __init__(self):
        self.msgs, self.ok = [], True

    def add(self, tag, cond, detail=""):
        if not cond:
            self.ok = False
        self.msgs.append("%s %s%s" % ("✓" if cond else "✗", tag, ("：" + detail) if detail else ""))

    def dump(self):
        print("\n".join("   " + m for m in self.msgs))
        return self.ok


# ─────────────────────────── fund ───────────────────────────

def plan_fund():
    wb, ws = load(F_REWARD)
    idx = header_index(ws)
    out = []
    for r, v in iter_data(ws, ROW_START[F_REWARD]):
        act = v[idx["ActivityId"]]
        if act not in FUND_SKU_ACT.values():
            continue
        reward = v[idx["Reward"]] or ""
        d = pairs_dict(reward)
        if DIA not in d:
            sys.exit("活动 %s 第 %d 行的 Reward 里没有钻石：%r" % (act, r, reward))
        out.append({"row": r, "act": act, "old": reward,
                    "old_dia": d[DIA], "keep": [k for k in d if k != DIA]})
    wb.close()
    if not out:
        sys.exit("没找到成长基金的奖励行，检查 ActivityId 是否被改过")
    return out


def apply_fund(plans):
    wb, ws = load(F_REWARD)
    idx = header_index(ws)
    c = idx["Reward"]
    for p in plans:
        base = pairs_dict(ws.cell(p["row"], c + 1).value)
        base[DIA] = FUND_NODE_DIA
        # 保持原有顺序，只改钻石数值
        order = [(k, (FUND_NODE_DIA if k == DIA else base[k])) for k in base]
        ws.cell(p["row"], c + 1).value = fmt_pairs(order)
    # TotalValues 只做展示（客户端 Lua 与服务端都不读），同步成真实总量免得虚标
    if "TotalValues" in idx:
        tv = idx["TotalValues"]
        for p in plans:
            if ws.cell(p["row"], tv + 1).value:
                ws.cell(p["row"], tv + 1).value = FUND_NODE_DIA * 10
    wb.save(path_of(F_REWARD))
    wb.close()


def verify_fund():
    ck = Checker()
    wb, ws = load(F_REWARD)
    idx = header_index(ws)
    per_act = {}
    for r, v in iter_data(ws, ROW_START[F_REWARD]):
        act = v[idx["ActivityId"]]
        if act not in FUND_SKU_ACT.values():
            continue
        d = pairs_dict(v[idx["Reward"]])
        per_act.setdefault(act, []).append((r, d))

    ck.add("五个成长基金活动的节点数都是 10",
           all(len(x) == 10 for x in per_act.values()) and len(per_act) == 5,
           str({k: len(v) for k, v in per_act.items()}))

    bad = [(r, d.get(DIA)) for lst in per_act.values() for r, d in lst
           if d.get(DIA) != FUND_NODE_DIA]
    ck.add("每个节点都是 %d 钻" % FUND_NODE_DIA, not bad, str(bad[:5]))

    extra = [(r, sorted(k for k in d if k != DIA)) for lst in per_act.values() for r, d in lst
             if any(k != DIA for k in d)]
    ck.add("节点里没有混入钻石以外的道具", not extra, str(extra[:4]))

    wb.close()

    # 抽/元：活动总量 / 98 元
    rates = {}
    for sku, act in FUND_SKU_ACT.items():
        total = FUND_NODE_DIA * 10
        rates[sku] = total / DIA_PER_PULL / FUND_PRICE
    ck.add("五个成长基金都是 1.000 抽/元",
           all(abs(x - 1.0) < 1e-6 for x in rates.values()),
           str({k: round(v, 4) for k, v in rates.items()}))

    # SKU 的 OnSaleRule 必须仍然指向同一个活动（防止改错行）
    wb2, ws2 = load(F_RECHARGE)
    idx2 = header_index(ws2)
    mismatch = []
    for r, v in iter_data(ws2, ROW_START[F_RECHARGE]):
        sid = v[idx2["Id"]]
        if sid not in FUND_SKU_ACT:
            continue
        rule = str(v[idx2["OnSaleRule"]] or "").replace(" ", "")
        if not rule.startswith("3#%d" % FUND_SKU_ACT[sid]):
            mismatch.append((sid, rule))
        if v[idx2["Price"]] != FUND_PRICE:
            mismatch.append((sid, "price=%s" % v[idx2["Price"]]))
    wb2.close()
    ck.add("SKU 101-105 仍指向活动 6/601-604 且价格未动", not mismatch, str(mismatch))
    return ck


# ─────────────────────────── pack ───────────────────────────

def _live_packs(ws, idx, billing):
    for r, v in iter_data(ws, ROW_START[F_RECHARGE]):
        sid = v[idx["Id"]]
        if sid is None or v[idx["Type"]] != PACK_TYPE:
            continue
        if not billing.get(sid, False):
            continue
        if str(v[idx["OnSaleRule"]] or "").replace(" ", "").startswith("1#999#999"):
            continue
        if (v[idx["Price"]] or 0) <= 0:
            continue
        yield r, v, sid


def _billing_map():
    wb, ws = load("ChannelRechargeIdConfig")
    out = {}
    for r, v in iter_data(ws, 8):
        out[v[1]] = any(x not in (None, "", 0) for x in v[3:])
    wb.close()
    return out


def plan_pack():
    """按价格档整体 ×1.21。

    两道闸门（都是粒度问题，不是设计选择）：
      1) 档位闸：整档中位抽数 < PACK_MIN_TIER_PULL 时整档不动 ——
         6 元档中位 1.55 抽，补 1 张就是 +65%，会把 6 元包顶到全表最甜，
         直接造出「小额比大额值」的反向阶梯。
      2) SKU 闸：算出来的实际增幅必须落在 MULT±TOL，否则这个 SKU 不动
         （同档里抽数极低的离群件补 1 张就 +50%，留到「削峰补谷」那一步处理）。
    """
    billing = _billing_map()
    wb, ws = load(F_RECHARGE)
    idx = header_index(ws)

    tiers = collections.defaultdict(list)
    for r, v, sid in _live_packs(ws, idx, billing):
        if sid in PACK_EXCLUDE:
            continue
        rewards = pairs_dict(v[idx["BaseReward"]])
        for k, n in pairs_dict(v[idx["ExtraReward"]]).items():
            rewards[k] = rewards.get(k, 0) + n
        tiers[v[idx["Price"]]].append({
            "row": r, "id": sid, "price": v[idx["Price"]],
            "rewards": rewards, "pull": pull_equiv(rewards),
            "base": v[idx["BaseReward"]] or "", "show": v[idx["RewardShow"]] or "",
        })
    wb.close()

    plans, skipped = [], collections.Counter()
    for price in sorted(tiers):
        rows = tiers[price]
        pulls = [x["pull"] for x in rows if x["pull"] > 0]
        empty = [x for x in rows if x["pull"] <= 0]
        if empty:
            skipped["纯资源包（无券无钻，乘不了）"] += len(empty)
        if not pulls:
            continue
        med = statistics.median(pulls)
        if med < PACK_MIN_TIER_PULL:
            skipped["整档不动（%s 元中位 %.2f 抽，补 1 张 = +%.0f%%）"
                    % (price, med, 100.0 / med)] += len(pulls)
            continue
        for x in rows:
            if x["pull"] <= 0:
                continue
            add = int(round(x["pull"] * (PACK_MULT - 1)))
            ratio = (x["pull"] + add) / x["pull"]
            if add < 1 or not (PACK_MULT - PACK_REL_TOL <= ratio <= PACK_MULT + PACK_REL_TOL):
                skipped["档内离群件（%s 元 / %.2f 抽，补 1 张 = +%.0f%%）"
                        % (price, x["pull"], 100.0 / x["pull"])] += 1
                continue
            plans.append(dict(x, ticket=pick_ticket(x["rewards"]), add=add,
                              pull_after=x["pull"] + add))
    return plans, skipped


def apply_pack(plans):
    wb, ws = load(F_RECHARGE)
    idx = header_index(ws)
    c_base, c_show = idx["BaseReward"], idx["RewardShow"]
    for p in plans:
        delta = {p["ticket"]: p["add"]}
        base = ws.cell(p["row"], c_base + 1).value or ""
        ws.cell(p["row"], c_base + 1).value = add_items(base, delta)
        show = ws.cell(p["row"], c_show + 1).value or ""
        if str(show).strip():
            ws.cell(p["row"], c_show + 1).value = add_items(show, delta)
    wb.save(path_of(F_RECHARGE))
    wb.close()


def verify_pack(plans_before):
    """plans_before: apply 之前算出的 plan（含 pull / add / price）。"""
    ck = Checker()
    billing = _billing_map()
    wb, ws = load(F_RECHARGE)
    idx = header_index(ws)
    cur = {}
    n_pack = 0
    for r, v, sid in _live_packs(ws, idx, billing):
        n_pack += 1
        rewards = pairs_dict(v[idx["BaseReward"]])
        for k, n in pairs_dict(v[idx["ExtraReward"]]).items():
            rewards[k] = rewards.get(k, 0) + n
        cur[sid] = {"pull": pull_equiv(rewards), "price": v[idx["Price"]],
                    "rewards": rewards, "base": v[idx["BaseReward"]],
                    "show": v[idx["RewardShow"]],
                    "ticket_tot": sum(rewards.get(k, 0) for k in TICKETS)}
    wb.close()

    ck.add("可售礼包数量与计划一致", n_pack >= len(plans_before), "%d" % n_pack)

    low, high, bad_dia, bad_other = [], [], [], []
    before = {p["id"]: p for p in plans_before}
    for sid, p in before.items():
        c = cur.get(sid)
        if c is None:
            low.append((sid, "missing"))
            continue
        ratio = c["pull"] / p["pull"] if p["pull"] else 0
        if not (PACK_MULT - PACK_REL_TOL - 1e-9 <= ratio <= PACK_MULT + PACK_REL_TOL + 1e-9):
            low.append((sid, round(ratio, 4)))
        if c["price"] != p["price"]:
            bad_dia.append((sid, p["price"], c["price"]))
    ck.add("每个被改的礼包增幅落在 %.2f±%.0f%%" % (PACK_MULT, PACK_REL_TOL * 100),
           not low, str(low[:5]))
    ck.add("礼包价格一个字没动", not bad_dia, str(bad_dia[:5]))

    # 钻石数量、非券道具构成必须与改前一致
    wb2, ws2 = load(F_RECHARGE)
    idx2 = header_index(ws2)
    drift = []
    for r, v, sid in _live_packs(ws2, idx2, billing):
        if sid not in before:
            continue
        now = pairs_dict(v[idx2["BaseReward"]])
        old = pairs_dict(before[sid]["base"] or "")
        if now.get(DIA, 0) != old.get(DIA, 0):
            drift.append((sid, DIA, old.get(DIA, 0), now.get(DIA, 0)))
        ok_keys = {k for k in old if k not in TICKETS}
        nk_keys = {k for k in now if k not in TICKETS}
        if ok_keys != nk_keys:
            drift.append((sid, "keys", sorted(ok_keys), sorted(nk_keys)))
    wb2.close()
    ck.add("钻石与非券道具构成零漂移", not drift, str(drift[:5]))

    # 奖励串不能出现非数字 id（fmt_pairs 出错的典型症状）
    illegal = [(sid, k) for sid, c in cur.items() for k in c["rewards"] if not k.isdigit()]
    ck.add("奖励串 id 全为数字", not illegal, str(illegal[:4]))

    # RewardShow 与 BaseReward 的券数量必须平行（一个给奖一个展示）
    diverge = []
    for sid, c in cur.items():
        show = str(c["show"] or "").strip()
        if not show:
            continue
        if sum(pairs_dict(show).get(k, 0) for k in TICKETS) != c["ticket_tot"]:
            diverge.append(sid)
    ck.add("RewardShow 券数量与 BaseReward 一致", not diverge, str(diverge[:5]))
    return ck


# ─────────────────────────── store ───────────────────────────

def _ticket_target(n):
    return STORE_TEN_UNIT * n if n >= 10 else STORE_UNIT * n


def plan_store():
    wb, ws = load(F_STORE)
    idx = header_index(ws)
    plans, already = [], 0
    c_goods, c_cost = idx["Goods"], idx["Cost"]
    for r, v in iter_data(ws, ROW_START[F_STORE]):
        goods = str(v[c_goods] or "").strip()
        cost = str(v[c_cost] or "").strip()
        if goods.count("|") or "#" not in goods:
            continue                                  # 组合包不动
        gid, _, gnum = goods.partition("#")
        if gid.strip() != TICKET_GOODS:
            continue
        try:
            n = int(float(gnum))
        except ValueError:
            continue
        head, _, tail = cost.partition("|")
        if head.strip() != DIA or "#" not in tail:
            continue                                  # 非钻石货币不动
        parts = tail.split("#")
        try:
            old = int(float(parts[-1]))
        except ValueError:
            continue
        new = _ticket_target(n)
        if new >= old:
            already += 1
            continue
        plans.append({"row": r, "id": v[idx["Id"]], "n": n,
                      "old": old, "new": new, "raw": cost})
    wb.close()
    return plans, already


def apply_store(plans):
    wb, ws = load(F_STORE)
    idx = header_index(ws)
    c = idx["Cost"]
    for p in plans:
        parts = p["raw"].split("|")
        tail = parts[-1].split("#")
        tail[-1] = str(p["new"])
        parts[-1] = "#".join(tail)
        ws.cell(p["row"], c + 1).value = "|".join(parts)
    wb.save(path_of(F_STORE))
    wb.close()


def verify_store(plans_before):
    ck = Checker()
    wb, ws = load(F_STORE)
    idx = header_index(ws)
    over, raised, kept = [], [], 0
    c_goods, c_cost = idx["Goods"], idx["Cost"]
    for r, v in iter_data(ws, ROW_START[F_STORE]):
        goods = str(v[c_goods] or "").strip()
        cost = str(v[c_cost] or "").strip()
        if goods.count("|") or "#" not in goods:
            continue
        gid, _, gnum = goods.partition("#")
        if gid.strip() != TICKET_GOODS:
            continue
        head, _, tail = cost.partition("|")
        if head.strip() != DIA or "#" not in tail:
            continue
        try:
            n = int(float(gnum))
            price = int(float(tail.split("#")[-1]))
        except ValueError:
            continue
        unit = price / float(n)
        if unit > STORE_UNIT + 1e-9:
            over.append((v[idx["Id"]], goods, price, round(unit, 1)))
        kept += 1
    wb.close()

    ck.add("卖券单价不再高于单抽 220 钻", not over, str(over[:6]))

    before = {p["id"]: p for p in plans_before}
    dry = load(F_STORE)[0]
    ws2 = dry[F_STORE] if F_STORE in dry.sheetnames else dry.worksheets[0]
    idx2 = header_index(ws2)
    for r, v in iter_data(ws2, ROW_START[F_STORE]):
        sid = v[idx2["Id"]]
        if sid not in before:
            continue
        p = before[sid]
        now = str(v[idx2["Cost"]] or "")
        if int(float(now.split("|")[-1].split("#")[-1])) > p["old"]:
            raised.append((sid, p["old"], now))
        if str(v[idx2["Goods"]]) != "%s#%d" % (TICKET_GOODS, p["n"]):
            raised.append((sid, "goods changed", v[idx2["Goods"]]))
    dry.close()
    ck.add("只降价、不涨价，且商品数量未变", not raised, str(raised[:6]))

    # 19#10 必须等于十连直抽价 2000；19#1 必须等于单抽钻石价 220
    probe = []
    wb3, ws3 = load(F_STORE)
    idx3 = header_index(ws3)
    for r, v in iter_data(ws3, ROW_START[F_STORE]):
        goods = str(v[idx3["Goods"]] or "").strip()
        cost = str(v[idx3["Cost"]] or "").strip()
        if goods in ("19#1", "19#10") and cost.startswith(DIA + "|"):
            n = 1 if goods == "19#1" else 10
            price = int(float(cost.split("|")[-1].split("#")[-1]))
            if price != _ticket_target(n):
                probe.append((v[idx3["Id"]], goods, price))
    wb3.close()
    ck.add("19#1=220 / 19#10=2000 全部对齐", not probe, str(probe[:6]))
    ck.add("扫到的钻石券行数", kept > 0, "%d 行" % kept)
    return ck


# ─────────────────────────── lux328 ───────────────────────────

def _fund_total(fund_type):
    """返回 (钻石合计, 券合计, 天数)。"""
    wb, ws = load(F_FUND)
    idx = header_index(ws)
    dia = tic = 0
    days = 0
    for r, v in iter_data(ws, ROW_START[F_FUND]):
        if v[idx["Type"]] != fund_type:
            continue
        days += 1
        d = pairs_dict(v[idx["reward"]])
        dia += d.get(DIA, 0)
        tic += sum(d.get(k, 0) for k in TICKETS)
    wb.close()
    return dia, tic, days


def plan_lux():
    wb, ws = load(F_FUND)
    idx = header_index(ws)
    plans = []
    for r, v in iter_data(ws, ROW_START[F_FUND]):
        if v[idx["Type"]] != LUX_TYPE_OLD:
            continue
        d = pairs_dict(v[idx["reward"]])
        if d.get(DIA) != LUX_DAY_OLD:
            continue
        plans.append({"row": r, "raw": v[idx["reward"]], "old": LUX_DAY_OLD})
    wb.close()
    # SKU 79 的 RewardShow 是「总共给多少钻」的展示镜像，要跟着改
    wb2, ws2 = load(F_RECHARGE)
    idx2 = header_index(ws2)
    show = None
    for r, v in iter_data(ws2, ROW_START[F_RECHARGE]):
        if v[idx2["Id"]] != LUX_SKU:
            continue
        if str(v[idx2["RewardShow"]] or "").strip() == LUX_SKU_SHOW:
            show = {"row": r, "old": LUX_SKU_SHOW}
    wb2.close()
    return plans, show


def apply_lux(plans, show):
    wb, ws = load(F_FUND)
    idx = header_index(ws)
    c = idx["reward"]
    for p in plans:
        d = pairs_dict(p["raw"])
        d[DIA] = LUX_DAY_NEW
        ws.cell(p["row"], c + 1).value = fmt_pairs([(k, d[k]) for k in d])
    wb.save(path_of(F_FUND))
    wb.close()
    if show:
        wb2, ws2 = load(F_RECHARGE)
        idx2 = header_index(ws2)
        new_total = LUX_DAY_NEW * len(plans) + (int(LUX_SKU_SHOW.split("#")[1]) - LUX_DAY_OLD * len(plans))
        ws2.cell(show["row"], idx2["RewardShow"] + 1).value = "%s#%d" % (DIA, new_total)
        wb2.save(path_of(F_RECHARGE))
        wb2.close()
        show["new"] = new_total


def verify_lux():
    ck = Checker()
    dia79, tic79, days79 = _fund_total(LUX_TYPE_OLD)
    dia78, tic78, _ = _fund_total(78)

    wb, ws = load(F_RECHARGE)
    idx = header_index(ws)
    base = {}
    for r, v in iter_data(ws, ROW_START[F_RECHARGE]):
        if v[idx["Id"]] in (78, 79):
            d = pairs_dict(v[idx["BaseReward"]])
            base[v[idx["Id"]]] = d.get(DIA, 0)
    wb.close()

    rate79 = (dia79 + base.get(79, 0)) / DIA_PER_PULL + tic79
    rate79 /= 328.0
    rate78 = (dia78 + base.get(78, 0)) / DIA_PER_PULL + tic78
    rate78 /= 128.0

    ck.add("Type=79 仍是 30 天", days79 == 30, str(days79))
    ck.add("328 元档 ≥ 128 元档（倒挂消除）", rate79 + 1e-9 >= rate78,
           "%.4f vs %.4f" % (rate79, rate78))
    ck.add("两档差距 ≤ 0.5%", abs(rate79 - rate78) <= 0.005 * rate78,
           "差 %.2f%%" % ((rate79 / rate78 - 1) * 100))
    ck.add("Type=78 未被波及", dia78 == 13200, str(dia78))

    # 展示镜像必须等于真实总量
    wb2, ws2 = load(F_RECHARGE)
    idx2 = header_index(ws2)
    bad = []
    for r, v in iter_data(ws2, ROW_START[F_RECHARGE]):
        if v[idx2["Id"]] in (78, 79):
            d = pairs_dict(v[idx2["RewardShow"]])
            real = base[v[idx2["Id"]]] + _fund_total(v[idx2["Id"]])[0]
            if d.get(DIA) != real:
                bad.append((v[idx2["Id"]], d.get(DIA), real))
    wb2.close()
    ck.add("SKU 78/79 的 RewardShow = 真实钻石总量", not bad, str(bad))
    return ck


# ─────────────────────────── 报告 ───────────────────────────

def print_pack_report(plans, skipped):
    if not plans:
        print("   （无需要补券的礼包）")
        return
    by = {}
    for p in plans:
        by.setdefault(p["price"], []).append(p)
    print("   %-6s %-5s %-9s %-9s %-6s %-8s" % ("价格", "件数", "改前中位", "改后中位", "补券", "实际倍数"))
    tot = 0
    for price in sorted(by):
        rows = by[price]
        b = statistics.median([x["pull"] / price for x in rows])
        a = statistics.median([x["pull_after"] / price for x in rows])
        add = sum(x["add"] for x in rows)
        tot += add
        print("   %-6s %-5d %-9.4f %-9.4f %-6d x%.3f"
              % (price, len(rows), b, a, add, a / b))
    print("   合计补券 %d 张 / %d 个礼包" % (tot, len(plans)))
    for reason, n in skipped.items():
        print("   跳过 %-4d 个：%s" % (n, reason))


def print_store_report(plans, already):
    if not plans:
        print("   （无需要改价的券）")
        return
    print("   %-6s %-8s %-7s %-7s %-8s" % ("行 id", "商品", "改前", "改后", "折扣"))
    for p in sorted(plans, key=lambda x: (x["n"], x["id"])):
        print("   %-6s %-8s %-7d %-7d -%.1f%%"
              % (p["id"], "19#%d" % p["n"], p["old"], p["new"],
                 (1 - p["new"] / float(p["old"])) * 100))
    print("   共 %d 行降价；已合规跳过 %d 行" % (len(plans), already))


def print_lux_report(plans, show):
    dia79, tic79, days = _fund_total(LUX_TYPE_OLD)
    new_dia = dia79 + (LUX_DAY_NEW - LUX_DAY_OLD) * len(plans)
    print("   Type=79 日领 %d → %d，%d 天" % (LUX_DAY_OLD, LUX_DAY_NEW, len(plans)))
    print("   30 天钻石 %d → %d（+%d），券 %d 张不变" % (dia79, new_dia, new_dia - dia79, tic79))
    if show:
        sku_base = int(LUX_SKU_SHOW.split("#")[1]) - LUX_DAY_OLD * len(plans)
        print("   SKU 79 的 RewardShow 同步 → 16#%d（= 底包 %d + 日领 %d）"
              % (new_dia + sku_base, sku_base, new_dia))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--force", action="store_true",
                    help="强制执行已经落过表的项（会叠加乘数，慎用）")
    ap.add_argument("--verify-only", action="store_true")
    ap.add_argument("--only", default="fund,pack,store,lux328",
                    help="逗号分隔：fund,pack,store,lux328")
    ap.add_argument("--revert", metavar="DIR")
    ap.add_argument("--list-backups", action="store_true")
    args = ap.parse_args()

    if args.list_backups:
        if not os.path.isdir(BACKUP_ROOT):
            print("没有备份")
            return 0
        for d in sorted(os.listdir(BACKUP_ROOT)):
            print(os.path.join(BACKUP_ROOT, d))
        return 0

    if args.revert:
        d = args.revert
        if not os.path.isdir(d):
            sys.exit("备份目录不存在：%s" % d)
        for f in os.listdir(d):
            shutil.copy2(os.path.join(d, f), os.path.join(BASE, f))
        print("已回滚 %d 个文件 ← %s" % (len(os.listdir(d)), d))
        return 0

    items = [x.strip() for x in args.only.split(",") if x.strip()]
    unknown = set(items) - {"fund", "pack", "store", "lux328"}
    if unknown:
        sys.exit("未知选项：%s" % unknown)

    print("=" * 62)
    print("R13 结构性定价修复  items=%s  mode=%s"
          % (",".join(items),
             "verify" if args.verify_only else ("apply" if args.apply else "dry-run")))
    print("=" * 62)

    if args.apply:
        guard(items, args.force)

    ok = True

    # ── fund ──────────────────────────────────────────────
    if "fund" in items:
        plans = plan_fund()
        print("\n[fund] 成长基金节点统一")
        by_act = {}
        for p in plans:
            by_act.setdefault(p["act"], []).append(p)
        for act in sorted(by_act):
            rows = by_act[act]
            old = sum(x["old_dia"] for x in rows)
            print("   活动 %-4s 节点 %2d 个  %5d 钻 → %5d 钻   %.3f → %.3f 抽/元"
                  % (act, len(rows), old, FUND_NODE_DIA * 10,
                     old / DIA_PER_PULL / FUND_PRICE,
                     FUND_NODE_DIA * 10 / DIA_PER_PULL / FUND_PRICE))
        if args.apply:
            bak = backup([path_of(F_REWARD)])
            print("   备份 → %s" % bak)
            apply_fund(plans)
            c = verify_fund()
            if not c.dump():
                _restore(bak)
                sys.exit("❌ fund 自检失败，已自动回滚")
            save_marker(["fund"], bak)
            ok &= c.ok
        elif args.verify_only:
            c = verify_fund()
            ok &= c.dump()

    # ── pack ──────────────────────────────────────────────
    plans_pack = []
    if "pack" in items:
        plans_pack, skipped = plan_pack()
        print("\n[pack] 礼包带 ×%.2f（只加券，钻石不动）" % PACK_MULT)
        print_pack_report(plans_pack, skipped)
        if args.apply and plans_pack:
            bak = backup([path_of(F_RECHARGE)])
            print("   备份 → %s" % bak)
            apply_pack(plans_pack)
            c = verify_pack(plans_pack)
            if not c.dump():
                _restore(bak)
                sys.exit("❌ pack 自检失败，已自动回滚")
            save_marker(["pack"], bak)
            ok &= c.ok
        elif not args.apply:
            print("   示例（前 6 个）：")
            for p in plans_pack[:6]:
                print("     id=%-6s %s 元  +%s#%d → %.3f 抽/元"
                      % (p["id"], p["price"], p["ticket"], p["add"],
                         p["pull_after"] / p["price"]))
        elif args.verify_only:
            print("   （verify-only 不做逐项比对，用 audit 脚本核对结果）")

    # ── store ─────────────────────────────────────────────
    if "store" in items:
        plans_store, already = plan_store()
        print("\n[store] 紫晶卡钻石标价 → 单抽 220 / 十连 2000")
        print_store_report(plans_store, already)
        if args.apply and plans_store:
            bak = backup([path_of(F_STORE)])
            print("   备份 → %s" % bak)
            apply_store(plans_store)
            c = verify_store(plans_store)
            if not c.dump():
                _restore(bak)
                sys.exit("❌ store 自检失败，已自动回滚")
            save_marker(["store"], bak)
            ok &= c.ok
        elif args.verify_only:
            c = verify_store([])
            ok &= c.dump()

    # ── lux328 ────────────────────────────────────────────
    if "lux328" in items:
        lux_plans, lux_show = plan_lux()
        print("\n[lux328] 豪华基金 328 元拉平")
        print_lux_report(lux_plans, lux_show)
        if args.apply:
            bak = backup([path_of(F_FUND), path_of(F_RECHARGE)])
            print("   备份 → %s" % bak)
            apply_lux(lux_plans, lux_show)
            c = verify_lux()
            if not c.dump():
                _restore(bak)
                sys.exit("❌ lux328 自检失败，已自动回滚")
            save_marker(["lux328"], bak)
            ok &= c.ok
        elif args.verify_only:
            c = verify_lux()
            ok &= c.dump()

    if not args.apply and not args.verify_only:
        print("\n[dry-run] 未写盘。确认无误后加 --apply。")
    elif args.apply:
        print("\n✅ 落表完成。回滚：--revert <backup_dir>（--list-backups 查看）")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
