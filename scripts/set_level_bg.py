#!/usr/bin/env python3
"""MainLevelConfig.BG 批量改写 —— 把某一章 / 某个 Id 区间的挂机战斗背景换成指定地图资源名。

地图名必须能在 ResourcePathConfig 里解析到真 prefab（resMgr:LoadAsset(name) 按名取资源），
所以先跑一遍登记校验，再动表。

用法：
  # 先看会改什么（不写文件）
  python3 scripts/set_level_bg.py --bg Map9 --chapter 15 --dry-run
  # 只改某段 Id
  python3 scripts/set_level_bg.py --bg Map9 --ids 15011-15405
  # 单关
  python3 scripts/set_level_bg.py --bg Map9 --ids 15405

改完必须导表（xlsx → 客户端 lua + LuaBytes），否则游戏里读的还是旧值：
  Jenkins → TS-Resource-Build（EXPORT_TABLES=true / SYNC_BYTES=true）
"""
from __future__ import annotations

import argparse
import os
import re
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "data_execl" / "base_data" / "MainLevelConfig.xlsx"
RES_CFG = ROOT / "Assets" / "ManagedResources" / "ResConfigs" / "ResourcePathConfig.asset"
PREFAB_DIR = ROOT / "Assets" / "ManagedResources" / "Prefabs" / "Battle"

COL_ID = "Id"
COL_BG = "BG"
FIRST_DATA_ROW = 8          # 1..7 是表头/类型/校对行


def resolve_map_name(name: str) -> tuple[bool, str]:
    """按 ResourcePathConfig 的口径校验地图资源名能不能被 resMgr 取到。"""
    try:
        import yaml
    except ImportError:
        yaml = None

    if yaml is not None and RES_CFG.exists():
        lines = RES_CFG.read_text(encoding="utf-8").split("\n")
        try:
            start = lines.index("MonoBehaviour:")
        except ValueError:
            return False, "ResourcePathConfig 结构异常（找不到 MonoBehaviour）"
        data = yaml.safe_load("\n".join(lines[start:]))["MonoBehaviour"]
        info_map = {x["resName"]: x for x in data["resInfoList"]}
        hit = info_map.get(name)
        if not hit:
            return False, f"ResourcePathConfig 里没有 {name}（真机 resMgr:LoadAsset 会返回 nil）"
        path = os.path.join(
            data["resPathList"][hit["resPathIndex"]], hit["resName"] + hit["extension"]
        )
        ab = data["resAbNameList"][hit["resAbNameIndex"]]
        if not os.path.exists(path):
            return False, f"登记了但文件不存在：{path}"
        return True, f"{path}  →  AB {ab}"

    prefab = PREFAB_DIR / f"{name}.prefab"
    if prefab.exists():
        return True, f"{prefab}（未校验 ResourcePathConfig，缺 pyyaml）"
    return False, f"找不到 {prefab}"


def parse_ids(spec: str) -> tuple[int, int]:
    spec = spec.strip()
    if "-" in spec:
        a, b = spec.split("-", 1)
        return int(a), int(b)
    v = int(spec)
    return v, v


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--bg", required=True, help="目标地图资源名，如 Map9")
    ap.add_argument("--chapter", type=int, help="整章：Id//1000 == 该值")
    ap.add_argument("--ids", help="Id 区间，如 15011-15405 或单值 15405")
    ap.add_argument("--dry-run", action="store_true", help="只报告不写文件")
    args = ap.parse_args()

    if not args.chapter and not args.ids:
        raise SystemExit("🔴 至少给 --chapter 或 --ids 之一")

    ok, msg = resolve_map_name(args.bg)
    print(("%s 资源校验通过：" % args.bg) + msg if ok else ("🔴 资源校验失败：%s" % msg))
    if not ok:
        return 2

    try:
        import openpyxl
    except ImportError:
        raise SystemExit("🔴 缺 openpyxl：pip install openpyxl")

    wb = openpyxl.load_workbook(XLSX)
    ws = wb[wb.sheetnames[0]]
    header = [c.value for c in ws[1]]
    col_id = header.index(COL_ID) + 1
    col_bg = header.index(COL_BG) + 1

    lo, hi = (0, 10**9)
    if args.ids:
        lo, hi = parse_ids(args.ids)
    ch = args.chapter

    changed: list[tuple[int, str, str]] = []
    for row in range(FIRST_DATA_ROW, ws.max_row + 1):
        lv = ws.cell(row=row, column=col_id).value
        if lv is None:
            continue
        try:
            lid = int(lv)
        except (TypeError, ValueError):
            continue
        if ch is not None and lid // 1000 != ch:
            continue
        if not (lo <= lid <= hi):
            continue
        old = ws.cell(row=row, column=col_bg).value
        if old == args.bg:
            continue
        changed.append((lid, str(old), args.bg))
        if not args.dry_run:
            ws.cell(row=row, column=col_bg).value = args.bg

    if not changed:
        print("没有需要改的行（已经是目标值或范围不匹配）")
        return 0

    before = Counter(c[1] for c in changed)
    print("将改 %d 行  Id %d..%d" % (len(changed), changed[0][0], changed[-1][0]))
    print("  原值分布: " + ", ".join("%s×%d" % (k, v) for k, v in before.most_common()))
    print("  首/末行: %s → %s / %s → %s" % (changed[0][0], args.bg, changed[-1][0], args.bg))

    if args.dry_run:
        print("（dry-run，未写文件）")
        return 0

    wb.save(XLSX)
    print("✓ 已写入 %s" % XLSX)
    print("下一步：Jenkins TS-Resource-Build（EXPORT_TABLES=true、SYNC_BYTES=true、REBUILD_AB 视需要）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
